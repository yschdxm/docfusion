from docx import Document
from typing import Dict, List, Any
import zipfile
import shutil
from .chunker import chunk_by_sections, chunk_table


class DocxParser:
    @staticmethod
    def parse(file_path: str) -> Dict[str, Any]:
        try:
            # 先检查文件是否是有效的ZIP文件
            if not zipfile.is_zipfile(file_path):
                raise ValueError("文件不是有效的docx格式（不是ZIP压缩包）")
            
            # 尝试用更宽容的方式打开docx
            try:
                doc = Document(file_path)
            except KeyError as e:
                # 如果缺少某些文件（如footnotes.xml），尝试用zipfile手动处理
                error_msg = str(e)
                if "footnotes" in error_msg.lower() or "endnotes" in error_msg.lower():
                    # 缺少脚注/尾注文件，尝试提取能读取的内容
                    doc = DocxParser._parse_with_missing_parts(file_path)
                else:
                    raise ValueError(f"文件损坏：docx文件结构不完整（缺少{error_msg}）")
        except zipfile.BadZipFile:
            raise ValueError("文件损坏：不是有效的docx文件（ZIP格式错误）")
        except Exception as e:
            if "文件损坏" in str(e) or "文件不是有效的" in str(e):
                raise e
            raise ValueError(f"文件解析失败：{str(e)}")
        
        content = []
        tables = []
        headings = []
        
        try:
            for para in doc.paragraphs:
                if para.text.strip():
                    content.append({
                        "text": para.text,
                        "style": para.style.name if para.style else None
                    })
                    if para.style and "Heading" in para.style.name:
                        headings.append({
                            "level": int(para.style.name[-1]) if para.style.name[-1].isdigit() else 1,
                            "text": para.text
                        })
        except Exception:
            # 如果解析段落失败，尝试用更简单的方式提取文本
            if not content:
                content = DocxParser._extract_text_from_zip(file_path)
        
        try:
            # 按文档顺序遍历，为每个表格提取上下文段落
            from docx.oxml.ns import qn
            table_contexts = []
            current_context = []

            for child in doc.element.body:
                if child.tag == qn('w:p'):
                    # 段落：收集为上下文
                    texts = []
                    for t in child.findall('.//' + qn('w:t')):
                        if t.text:
                            texts.append(t.text)
                    para_text = ''.join(texts).strip()
                    if para_text:
                        current_context.append(para_text)
                elif child.tag == qn('w:tbl'):
                    # 表格：提取数据，附带上文
                    table_data = []
                    for row in child.findall(qn('w:tr')):
                        row_data = []
                        for cell in row.findall(qn('w:tc')):
                            cell_texts = []
                            for t in cell.findall('.//' + qn('w:t')):
                                if t.text:
                                    cell_texts.append(t.text)
                            row_data.append(''.join(cell_texts))
                        table_data.append(row_data)
                    tables.append(table_data)
                    table_contexts.append("\n".join(current_context))
                    current_context = []  # 重置上下文
        except Exception:
            # 回退：简单遍历（无上下文）
            table_contexts = []
            try:
                for table in doc.tables:
                    table_data = []
                    for row in table.rows:
                        row_data = [cell.text for cell in row.cells]
                        table_data.append(row_data)
                    tables.append(table_data)
            except Exception:
                pass
        
        full_text = "\n".join([p["text"] for p in content])

        # 生成结构化 chunks
        chunks = chunk_by_sections(content, headings)
        chunk_idx = len(chunks)
        for table_data in tables:
            table_chunks = chunk_table(table_data)
            for tc in table_chunks:
                tc["chunk_index"] = chunk_idx
                chunk_idx += 1
            chunks.extend(table_chunks)

        return {
            "full_text": full_text,
            "paragraphs": content,
            "headings": headings,
            "tables": tables,
            "table_contexts": table_contexts,
            "chunks": chunks,
        }
    
    @staticmethod
    def _parse_with_missing_parts(file_path: str) -> Document:
        """尝试用更宽容的方式解析docx文件"""
        import tempfile
        import os
        
        # 创建临时目录
        temp_dir = tempfile.mkdtemp()
        temp_file = os.path.join(temp_dir, "temp.docx")
        
        try:
            # 复制docx文件并添加缺失的文件
            with zipfile.ZipFile(file_path, 'r') as source_zip:
                with zipfile.ZipFile(temp_file, 'w') as target_zip:
                    for item in source_zip.infolist():
                        data = source_zip.read(item.filename)
                        target_zip.writestr(item, data)
                    
                    # 添加缺失的空文件
                    empty_xml = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
                    empty_endnotes = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:endnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
                    
                    if 'word/footnotes.xml' not in source_zip.namelist():
                        target_zip.writestr('word/footnotes.xml', empty_xml)
                    if 'word/endnotes.xml' not in source_zip.namelist():
                        target_zip.writestr('word/endnotes.xml', empty_endnotes)
            
            return Document(temp_file)
        finally:
            # 清理临时文件
            shutil.rmtree(temp_dir, ignore_errors=True)
    
    @staticmethod
    def _extract_text_from_zip(file_path: str) -> List[Dict[str, Any]]:
        """直接从ZIP中提取文本内容"""
        content = []
        try:
            with zipfile.ZipFile(file_path, 'r') as zip_file:
                if 'word/document.xml' in zip_file.namelist():
                    import xml.etree.ElementTree as ET
                    xml_content = zip_file.read('word/document.xml')
                    root = ET.fromstring(xml_content)
                    
                    # 提取所有文本节点
                    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                    for para in root.findall('.//w:p', ns):
                        texts = []
                        for text_elem in para.findall('.//w:t', ns):
                            if text_elem.text:
                                texts.append(text_elem.text)
                        if texts:
                            content.append({
                                "text": ''.join(texts),
                                "style": None
                            })
        except Exception:
            pass
        return content
    
    @staticmethod
    def write(content: str, output_path: str, format_info: Dict = None):
        doc = Document()
        
        if format_info and "headings" in format_info:
            for heading in format_info["headings"]:
                doc.add_heading(heading["text"], level=heading.get("level", 1))
        
        for para_text in content.split("\n"):
            if para_text.strip():
                doc.add_paragraph(para_text)
        
        if format_info and "tables" in format_info:
            for table_data in format_info["tables"]:
                if table_data:
                    rows_count = len(table_data)
                    cols_count = len(table_data[0]) if table_data else 0
                    if rows_count > 0 and cols_count > 0:
                        table = doc.add_table(rows=rows_count, cols=cols_count)
                        for i, row_data in enumerate(table_data):
                            for j, cell_text in enumerate(row_data):
                                if j < len(table.rows[i].cells):
                                    table.rows[i].cells[j].text = str(cell_text) if cell_text else ""
        
        doc.save(output_path)

