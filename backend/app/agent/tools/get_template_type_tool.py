"""
获取模板类型工具 - 快速判断模板是数据表格还是表单

功能：
- 快速分析模板结构
- 返回 structure_type: data_table / form / mixed
- 用于路由决策，不进行详细字段检测
"""

import re
from typing import Any, Dict, Optional

from app.agent.base.tool import BaseTool, ToolContext, ToolResult
from app.db.postgres import async_session
from app.models.document import Document
from sqlalchemy import select


class GetTemplateTypeTool(BaseTool):
    """获取模板类型工具

    快速判断模板是数据表格还是表单，用于路由决策。
    比 get_form_structure 更轻量，只做类型判断，不检测具体字段。
    """

    @property
    def name(self) -> str:
        return "get_template_type"

    @property
    def description(self) -> str:
        return """快速判断模板是数据表格还是表单，用于决定使用哪个Agent处理。

返回 structure_type：
- "data_table": 数据表格，有表头行，下面有多行结构相同的数据行（如城市数据表、人员名单）
- "form": 表单，有多个不同标签的字段（如报名表、申请表、合同）
- "mixed": 混合类型

使用场景：
- 用户提供了template_id，需要判断是委派给 fill_table 还是 fill_form
- 快速路由决策，不需要检测具体字段"""

    @property
    def parameters(self) -> Dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "template_id": {
                    "type": "string",
                    "description": "模板文档ID"
                }
            },
            "required": ["template_id"]
        }

    async def execute(self, params: Dict[str, Any], context: ToolContext) -> ToolResult:
        """执行模板类型检测"""
        try:
            template_id = params.get("template_id", "")

            # UUID格式校验
            uuid_error = BaseTool.validate_uuid(template_id, "template_id")
            if uuid_error:
                return ToolResult(success=False, error=uuid_error)

            # 查询文档
            async with async_session() as db:
                result = await db.execute(
                    select(Document).where(Document.id == template_id)
                )
                doc = result.scalar_one_or_none()

                if not doc:
                    return ToolResult(
                        success=False,
                        error=f"模板文档不存在: {template_id}"
                    )

                # 支持 docx 和 xlsx 格式
                if doc.file_type not in ("docx", "xlsx"):
                    return ToolResult(
                        success=False,
                        error=f"模板类型检测支持docx和xlsx格式，当前格式: {doc.file_type}"
                    )

                # 快速检测模板类型
                structure_type = await self._detect_template_type(doc.file_path, doc.file_type)

                return ToolResult(
                    success=True,
                    data={
                        "template_id": template_id,
                        "filename": doc.original_filename,
                        "file_type": doc.file_type,
                        "structure_type": structure_type
                    }
                )

        except Exception as e:
            return ToolResult(
                success=False,
                error=f"模板类型检测失败: {str(e)}"
            )

    async def _detect_template_type(self, file_path: str, file_type: str) -> str:
        """快速检测模板类型"""
        from app.services.llm_service import llm_service

        # 提取模板内容（限制长度，提高速度）
        if file_type == "xlsx":
            content = self._extract_xlsx_summary(file_path)
        else:
            content = self._extract_docx_summary(file_path)

        # 使用LLM快速判断类型
        prompt = f"""请快速判断以下模板是数据表格还是表单。

模板内容：
{content}

只返回以下之一（不要有其他内容）：
- data_table: 数据表格，有表头行，下面有多行结构相同的数据行需要填写
- form: 表单，有多个不同标签的字段需要填写
- mixed: 混合类型"""

        try:
            response = await llm_service.chat_completion(
                messages=[{"role": "user", "content": prompt}],
                temperature=0.1,
                enable_thinking=False,
                max_tokens=50  # 只需要简短回答
            )

            # 解析响应
            if isinstance(response, str):
                response = response.strip().lower()
                if "data_table" in response:
                    return "data_table"
                elif "form" in response:
                    return "form"
                elif "mixed" in response:
                    return "mixed"

            return "unknown"

        except Exception as e:
            print(f"LLM类型检测失败: {e}")
            return "unknown"

    def _extract_docx_summary(self, file_path: str) -> str:
        """提取docx文件摘要（前几行和表格结构）"""
        from docx import Document as DocxDocument

        doc = DocxDocument(file_path)
        summary_parts = []

        # 提取前10个段落
        for i, para in enumerate(doc.paragraphs[:10]):
            text = para.text.strip()
            if text:
                summary_parts.append(f"[段落{i}] {text[:100]}")

        # 提取表格结构（前3行）
        for table_idx, table in enumerate(doc.tables[:3]):
            summary_parts.append(f"\n[表格{table_idx}] {len(table.rows)}行 x {len(table.columns)}列")
            for row_idx, row in enumerate(table.rows[:3]):
                cells = [cell.text.strip()[:30] for cell in row.cells]
                summary_parts.append(f"  行{row_idx}: {' | '.join(cells)}")

        summary = "\n".join(summary_parts)
        if len(summary) > 2000:
            summary = summary[:2000] + "\n..."

        return summary

    def _extract_xlsx_summary(self, file_path: str) -> str:
        """提取xlsx文件摘要（表头和前几行数据）"""
        from openpyxl import load_workbook

        summary_parts = []
        wb = load_workbook(file_path, read_only=True, data_only=True)

        for sheet_name in wb.sheetnames[:3]:  # 最多3个工作表
            sheet = wb[sheet_name]
            summary_parts.append(f"\n[工作表: {sheet_name}]")

            for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_idx >= 5:  # 最多5行
                    break
                cells = [str(cell)[:30] if cell is not None else "(空)" for cell in row]
                summary_parts.append(f"  行{row_idx}: {' | '.join(cells)}")

        wb.close()

        summary = "\n".join(summary_parts)
        if len(summary) > 2000:
            summary = summary[:2000] + "\n..."

        return summary
