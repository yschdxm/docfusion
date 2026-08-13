"""
表单填写工具 - 填写Word/Excel文档中的表单字段

功能：
- 按 field_id 精确填写（主路径，field_id 来自 get_template_structure）
- 兼容 label 键值对自动匹配
- 检测并填写下划线/方括号/冒号占位符、表格表单单元格、Word内容控件
- 填写 xlsx 纵向表单（A列标签 B列值，按坐标写入）
"""

import logging
import os
import re
from difflib import SequenceMatcher
from typing import Any, Dict, List

from sqlalchemy import select

from app.agent.base.tool import BaseTool, ToolContext, ToolResult
from app.agent.tools.get_template_structure_tool import analyze_template
from app.db.postgres import async_session
from app.models.document import Document
from app.services import document_versioning, file_storage

logger = logging.getLogger(__name__)


class FillFormTool(BaseTool):
    """表单填写工具

    使用提供的键值对数据填写Word文档中的非结构化表单。
    支持下划线、方括号、冒号模式、表格表单和内容控件。
    """

    @property
    def name(self) -> str:
        return "fill_form"

    @property
    def description(self) -> str:
        return """使用提供的数据填写文档中的表单字段（Word表单和Excel纵向表单）。

前置步骤（必须）：
- 先调用 get_template_structure 获取字段列表（含稳定 field_id），
  然后用 fields 参数按 field_id 精确填写

参数说明（两种模式，优先用 fields）：
- fields（推荐）: 字段数组，每项 {"field_id": "F1", "value": "填写内容"}
  - field_id 来自 get_template_structure 返回的 fields[].field_id
  - 也接受 {"label": "姓名", "value": "张三"} 按标签匹配
- data（兼容模式）: 键值对 {"姓名": "张三"}，按标签自动匹配字段

支持的表单类型：
- Word: 下划线______、方括号【】、冒号模式（姓名：）、表格表单、内容控件
- Excel纵向表单: A列标签B列填值，按 get_template_structure 返回的坐标写入

填写模式：
- overwrite: 覆盖现有内容（首次填写使用）
- append: 追加到已有内容后面（增量填写使用）

增量填写：
- 首次填写传 template_id，返回 output_file_id
- 后续补充传 output_doc_id=上次返回的ID + template_id（用于字段定位）
- 未匹配的字段会在 unmatched_keys 中返回，请检查标签拼写"""

    @property
    def parameters(self) -> Dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "template_id": {
                    "type": "string",
                    "description": "模板文档ID（首次填写时必需；增量填写时也需提供，用于字段定位）"
                },
                "output_doc_id": {
                    "type": "string",
                    "description": "已生成的输出文档ID（增量填写时提供）"
                },
                "fields": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "field_id": {"type": "string", "description": "get_template_structure 返回的字段ID，如 F1"},
                            "label": {"type": "string", "description": "字段标签（未提供field_id时按标签匹配）"},
                            "value": {"description": "填写内容"}
                        }
                    },
                    "description": "要填写的字段数组（推荐）。每项提供 field_id 或 label + value"
                },
                "data": {
                    "type": "object",
                    "description": "兼容模式：填表数据键值对。key为字段标签（如\"姓名\"），value为填写内容",
                    "additionalProperties": True
                },
                "fill_mode": {
                    "type": "string",
                    "enum": ["overwrite", "append"],
                    "default": "overwrite",
                    "description": "填写模式：overwrite=覆盖现有内容，append=追加到已有内容后面"
                },
                "field_mapping": {
                    "type": "object",
                    "description": "手动字段映射（可选，仅 data 模式），当自动匹配失败时使用",
                    "additionalProperties": True
                }
            },
            "required": []
        }

    async def execute(self, params: Dict[str, Any], context: ToolContext) -> ToolResult:
        """执行表单填写"""
        try:
            template_id = params.get("template_id", "")
            output_doc_id = params.get("output_doc_id", "")
            field_items = params.get("fields") or []
            data = params.get("data") or {}
            fill_mode = params.get("fill_mode", "overwrite")
            field_mapping = params.get("field_mapping", {})

            if not field_items and not data:
                return ToolResult(success=False, error="填写内容不能为空（提供 fields 数组或 data 键值对）")

            # UUID校验
            if template_id:
                uuid_error = BaseTool.validate_uuid(template_id, "template_id")
                if uuid_error:
                    return ToolResult(success=False, error=uuid_error)
            if output_doc_id:
                uuid_error = BaseTool.validate_uuid(output_doc_id, "output_doc_id")
                if uuid_error:
                    return ToolResult(success=False, error=uuid_error)

            async with async_session() as db:
                run_id = context.metadata.get("run_id")
                conversation_id = context.metadata.get("conversation_id")

                # 确定操作的文件
                if output_doc_id:
                    # 增量填写：同 run 原地改；新 run 首次修改创建 version+1（文件已复制好）
                    result = await db.execute(
                        select(Document).where(Document.id == output_doc_id)
                    )
                    output_doc = result.scalar_one_or_none()
                    if not output_doc:
                        return ToolResult(success=False, error=f"输出文档不存在: {output_doc_id}")

                    output_doc, _is_new_version = await document_versioning.resolve_output_target(
                        db, output_doc,
                        run_id=run_id,
                        origin_type="fill",
                        conversation_id=conversation_id,
                    )
                    file_path = output_doc.file_path
                elif template_id:
                    # 首次填写：基于模板创建新 root 的输出 v1（文件复制自模板）
                    result = await db.execute(
                        select(Document).where(Document.id == template_id)
                    )
                    template_doc = result.scalar_one_or_none()
                    if not template_doc:
                        return ToolResult(success=False, error=f"模板文档不存在: {template_id}")

                    if template_doc.file_type not in ("docx", "xlsx"):
                        return ToolResult(
                            success=False,
                            error=f"表单填写支持docx/xlsx格式，当前格式: {template_doc.file_type}"
                        )

                    output_doc = await document_versioning.create_root_output(
                        db,
                        user_id=context.user_id,
                        file_type=template_doc.file_type,
                        origin_type="fill",
                        source_doc=template_doc,
                        run_id=run_id,
                        conversation_id=conversation_id,
                    )
                    file_path = output_doc.file_path
                else:
                    return ToolResult(success=False, error="首次填写必须提供template_id")

                # 获取表单结构：优先分析模板（结构与输出一致、缓存稳定命中）；
                # 未提供 template_id 时分析输出文档自身
                structure_doc_id = template_id or output_doc_id
                result = await db.execute(
                    select(Document).where(Document.id == structure_doc_id)
                )
                structure_doc = result.scalar_one_or_none()
                if not structure_doc:
                    return ToolResult(success=False, error=f"结构分析目标文档不存在: {structure_doc_id}")

                analysis = await analyze_template(structure_doc)
                form_fields = analysis.get("fields", [])
                if not form_fields:
                    return ToolResult(
                        success=False,
                        error="未检测到可填写的表单字段（structure_type="
                              f"{analysis.get('structure_type')}）。"
                              "若是数据表格请改用 fill_table"
                    )

                # 匹配数据到表单字段
                if field_items:
                    matched_fields, unmatched_keys = self._resolve_field_items(
                        field_items, form_fields
                    )
                else:
                    matched_fields = self._match_data_to_fields(
                        data, form_fields, field_mapping
                    )
                    unmatched_keys = list(set(data.keys()) - {m['data_key'] for m in matched_fields})

                # 执行填写
                if output_doc.file_type == "xlsx":
                    filled_count = self._fill_xlsx_fields(file_path, matched_fields)
                else:
                    filled_count = await self._fill_form_fields(
                        file_path, matched_fields, fill_mode
                    )

                # 更新输出版本的大小/哈希
                output_doc.file_size = os.path.getsize(file_path)
                output_doc.sha256 = file_storage.sha256_file(file_path)
                await db.commit()
                output_file_id = str(output_doc.id)

                return ToolResult(
                    success=True,
                    data={
                        "filled_fields": filled_count,
                        "total_fields": len(form_fields),
                        "matched_fields": len(matched_fields),
                        "unmatched_keys": unmatched_keys,
                        "output_file_id": output_file_id,
                        "output_filename": output_doc.original_filename,
                        "download_url": f"/api/v1/documents/{output_file_id}/download",
                        "field_details": [
                            {
                                "label": m['field_label'],
                                "data_key": m['data_key'],
                                "value": str(m['value'])[:50] + "..." if len(str(m['value'])) > 50 else str(m['value']),
                                "match_score": m['match_score']
                            }
                            for m in matched_fields
                        ]
                    },
                    metadata={
                        "template_id": template_id,
                        "fill_mode": fill_mode,
                        "is_new_file": not bool(output_doc_id)
                    }
                )

        except Exception as e:
            logger.exception(f"表单填写失败: {e}")
            return ToolResult(
                success=False,
                error=f"表单填写失败: {str(e)}"
            )

    def _resolve_field_items(
        self,
        field_items: List[Dict[str, Any]],
        form_fields: List[Dict[str, Any]],
    ) -> tuple:
        """解析 fields 数组：按 field_id 精确匹配，无 field_id 时按 label 匹配

        Returns:
            (matched_fields, unmatched_keys)
        """
        by_id = {f.get("field_id"): f for f in form_fields if f.get("field_id")}
        matched = []
        unmatched = []
        used_field_ids = set()

        for i, item in enumerate(field_items):
            value = item.get("value")
            field = None
            data_key = None

            field_id = item.get("field_id")
            if field_id:
                data_key = field_id
                field = by_id.get(field_id)
                if field is None:
                    unmatched.append(f"{field_id}(field_id不存在)")
                    continue
            else:
                label = item.get("label", "")
                data_key = label or f"第{i + 1}项"
                if not label:
                    unmatched.append(f"第{i + 1}项(缺少field_id和label)")
                    continue
                # 精确匹配优先，其次模糊匹配
                best_score = 0.0
                for f in form_fields:
                    if f.get("field_id") in used_field_ids:
                        continue
                    if f["label"] == label:
                        field = f
                        best_score = 1.0
                        break
                    score = self._calculate_match_score(label, f["label"])
                    if score > best_score:
                        best_score = score
                        field = f
                if field is None or best_score < 0.4:
                    unmatched.append(label)
                    continue

            used_field_ids.add(field.get("field_id"))
            matched.append({
                'field': field,
                'data_key': data_key,
                'field_label': field['label'],
                'value': value,
                'match_score': 1.0 if field_id else 0.9,
            })

        return matched, unmatched

    def _fill_xlsx_fields(self, file_path: str, matched_fields: List[Dict[str, Any]]) -> int:
        """填写 xlsx 纵向表单：按字段 location 的 sheet+cell 坐标写入"""
        from openpyxl import load_workbook

        wb = load_workbook(file_path)
        filled_count = 0

        for match in matched_fields:
            field = match['field']
            location = field.get('location', {})
            sheet_name = location.get('sheet')
            cell_coord = location.get('cell')
            if not sheet_name or not cell_coord:
                logger.warning(f"[FillFormTool] 字段 {field.get('label')} 缺少 sheet/cell 坐标，跳过")
                continue
            if sheet_name not in wb.sheetnames:
                logger.warning(f"[FillFormTool] 工作表不存在: {sheet_name}，跳过字段 {field.get('label')}")
                continue

            value = match['value']
            if isinstance(value, (list, dict)):
                value = "、".join(str(v) for v in value) if isinstance(value, list) else str(value)

            ws = wb[sheet_name]
            cell = ws[cell_coord]
            # 数值类型字段尝试按数值写入，避免Excel中数字变文本
            if field.get('detected_type') == 'number' and value is not None:
                try:
                    f = float(str(value).replace(',', ''))
                    value = int(f) if f.is_integer() else f
                except (TypeError, ValueError):
                    pass
            cell.value = value
            filled_count += 1

        wb.save(file_path)
        wb.close()
        return filled_count

    def _match_data_to_fields(
        self,
        data: Dict[str, Any],
        form_fields: List[Dict[str, Any]],
        manual_mapping: Dict[str, str]
    ) -> List[Dict[str, Any]]:
        """将数据匹配到表单字段"""
        matched = []
        used_field_indexes = set()

        # 先处理手动映射
        for data_key, field_label in manual_mapping.items():
            if data_key not in data:
                continue
            for idx, field in enumerate(form_fields):
                if idx in used_field_indexes:
                    continue
                if field['label'] == field_label:
                    matched.append({
                        'field': field,
                        'data_key': data_key,
                        'field_label': field_label,
                        'value': data[data_key],
                        'match_score': 1.0
                    })
                    used_field_indexes.add(idx)

        # 自动匹配剩余字段
        for data_key, value in data.items():
            # 跳过已手动映射的
            if any(m['data_key'] == data_key for m in matched):
                continue

            # 查找所有匹配的字段（处理重复标签）
            matching_fields = []
            for idx, field in enumerate(form_fields):
                if idx in used_field_indexes:
                    continue

                score = self._calculate_match_score(data_key, field['label'])
                if score >= 0.4:
                    matching_fields.append((idx, field, score))

            if not matching_fields:
                continue

            # 按匹配分数排序
            matching_fields.sort(key=lambda x: x[2], reverse=True)

            # 如果有多个字段匹配同一标签，全部填充
            # 这处理了"组名"需要填充多个单元格的情况
            best_score = matching_fields[0][2]
            for idx, field, score in matching_fields:
                # 只填充分数接近最佳分数的字段
                if score >= best_score * 0.9:
                    matched.append({
                        'field': field,
                        'data_key': data_key,
                        'field_label': field['label'],
                        'value': value,
                        'match_score': score
                    })
                    used_field_indexes.add(idx)

        return matched

    def _calculate_match_score(self, data_key: str, field_label: str) -> float:
        """计算数据key与字段标签的匹配分数"""
        # 精确匹配
        if data_key == field_label:
            return 1.0

        # 大小写不敏感匹配
        if data_key.lower() == field_label.lower():
            return 0.95

        # 包含匹配
        if data_key in field_label or field_label in data_key:
            return 0.8

        # 同义词匹配
        synonyms = {
            '姓名': ['名字', 'name', '用户名', '联系人'],
            '电话': ['手机', '联系方式', 'phone', 'tel', '联系电话'],
            '邮箱': ['邮件', 'email', '电子邮箱'],
            '地址': ['住址', 'address', '通讯地址'],
            '日期': ['时间', 'date', '签约日期'],
            '甲方': ['委托方', '客户方', '买方'],
            '乙方': ['受托方', '服务方', '卖方'],
        }

        for canonical, syn_list in synonyms.items():
            if (data_key in [canonical] + syn_list and field_label in [canonical] + syn_list):
                return 0.85

        # 序列匹配（模糊匹配）
        ratio = SequenceMatcher(None, data_key.lower(), field_label.lower()).ratio()
        return ratio

    async def _fill_form_fields(
        self,
        file_path: str,
        matched_fields: List[Dict[str, Any]],
        fill_mode: str
    ) -> int:
        """执行实际的表单填写"""
        from docx import Document as DocxDocument

        doc = DocxDocument(file_path)
        filled_count = 0

        # 按字段类型分组处理
        # 1. 处理普通字段（单值）
        # 2. 处理列表字段（多行数据，如成员列表）
        list_fields = []
        single_fields = []

        for match in matched_fields:
            value = match['value']
            if isinstance(value, list) and len(value) > 0 and isinstance(value[0], dict):
                list_fields.append(match)
            else:
                single_fields.append(match)

        # 处理单值字段
        for match in single_fields:
            field = match['field']
            value = str(match['value']) if match['value'] is not None else ""
            location = field['location']
            placeholder_type = field['placeholder_type']

            try:
                if placeholder_type == 'underline':
                    success = self._fill_underline(doc, location, value, fill_mode)
                elif placeholder_type == 'bracket':
                    success = self._fill_bracket(doc, location, value, fill_mode)
                elif placeholder_type == 'colon':
                    success = self._fill_colon(doc, location, value, fill_mode)
                elif placeholder_type == 'table_cell':
                    success = self._fill_table_cell(doc, location, value)
                elif placeholder_type == 'content_control':
                    success = self._fill_content_control(doc, location, value)
                elif placeholder_type == 'llm_detected':
                    # LLM检测到的字段，根据位置判断是段落还是表格
                    if 'paragraph_index' in location:
                        success = self._fill_paragraph_field(doc, location, value, fill_mode)
                    elif 'table_index' in location:
                        success = self._fill_table_cell(doc, location, value)
                    else:
                        success = False
                else:
                    success = False

                if success:
                    filled_count += 1
                    logger.info(f"已填写字段: {field['label']} = {value[:30]}...")
                else:
                    logger.warning(f"填写字段失败: {field['label']}")

            except Exception as e:
                logger.error(f"填写字段 {field['label']} 时出错: {e}")

        # 处理列表字段（如成员列表）
        for match in list_fields:
            field = match['field']
            values = match['value']  # List[Dict[str, str]]
            location = field['location']

            try:
                if location.get('is_header_based'):
                    # 表头行模式：填写多行数据
                    success = self._fill_table_cell_list(doc, location, values)
                    if success:
                        filled_count += len(values)
                        logger.info(f"已填写列表字段: {field['label']}, {len(values)} 行")
                    else:
                        logger.warning(f"填写列表字段失败: {field['label']}")
                else:
                    # 普通模式：只填写第一个值
                    if values:
                        first_value = str(values[0]) if not isinstance(values[0], dict) else str(list(values[0].values())[0])
                        success = self._fill_table_cell(doc, location, first_value)
                        if success:
                            filled_count += 1
                            logger.info(f"已填写字段: {field['label']} = {first_value[:30]}...")

            except Exception as e:
                logger.error(f"填写字段 {field['label']} 时出错: {e}")

        doc.save(file_path)
        return filled_count

    def _fill_underline(self, doc, location: Dict, value: str, fill_mode: str) -> bool:
        """替换下划线占位符"""
        para_idx = location.get('paragraph_index')
        if para_idx is None or para_idx >= len(doc.paragraphs):
            return False

        para = doc.paragraphs[para_idx]
        text = para.text

        # 找到下划线占位符
        match = re.search(r'_{3,}', text)
        if not match:
            return False

        if fill_mode == 'append':
            # 追加模式：在下划线前插入值
            new_text = text[:match.start()] + value + " " + text[match.start():]
        else:
            # 覆盖模式：替换下划线为值（保留部分下划线作为视觉提示）
            # 保留足够长度的下划线以维持格式
            underline_len = len(match.group())
            value_len = len(value)
            if value_len >= underline_len:
                new_text = text[:match.start()] + value + text[match.end():]
            else:
                # 值较短时，保留部分下划线
                remaining = '_' * (underline_len - value_len)
                new_text = text[:match.start()] + value + remaining + text[match.end():]

        # 保留格式：更新第一个run的文本
        if para.runs:
            para.runs[0].text = new_text
            for run in para.runs[1:]:
                run.text = ""
        else:
            para.text = new_text

        return True

    def _fill_bracket(self, doc, location: Dict, value: str, fill_mode: str) -> bool:
        """替换方括号占位符"""
        para_idx = location.get('paragraph_index')
        if para_idx is None or para_idx >= len(doc.paragraphs):
            return False

        para = doc.paragraphs[para_idx]
        text = para.text

        # 匹配方括号占位符
        patterns = [
            (re.compile(r'【\s*】'), '【', '】'),
            (re.compile(r'\[\s*\]'), '[', ']'),
            (re.compile(r'（\s*）'), '（', '）'),
        ]

        for pattern, open_bracket, close_bracket in patterns:
            match = pattern.search(text)
            if match:
                if fill_mode == 'append':
                    new_text = text[:match.start()] + open_bracket + value + close_bracket + " " + text[match.end():]
                else:
                    new_text = text[:match.start()] + open_bracket + value + close_bracket + text[match.end():]

                if para.runs:
                    para.runs[0].text = new_text
                    for run in para.runs[1:]:
                        run.text = ""
                else:
                    para.text = new_text
                return True

        return False

    def _fill_colon(self, doc, location: Dict, value: str, fill_mode: str) -> bool:
        """在冒号后追加值"""
        para_idx = location.get('paragraph_index')
        if para_idx is None or para_idx >= len(doc.paragraphs):
            return False

        para = doc.paragraphs[para_idx]
        text = para.text

        # 匹配行尾冒号
        match = re.search(r'[：:]\s*$', text)
        if not match:
            return False

        if fill_mode == 'append':
            new_text = text + value
        else:
            new_text = text.rstrip() + value

        if para.runs:
            para.runs[0].text = new_text
            for run in para.runs[1:]:
                run.text = ""
        else:
            para.text = new_text

        return True

    def _fill_table_cell(self, doc, location: Dict, value: str) -> bool:
        """填写表格单元格"""
        table_idx = location.get('table_index')
        row_idx = location.get('row_index')
        cell_idx = location.get('cell_index')

        if table_idx is None or row_idx is None or cell_idx is None:
            return False

        if table_idx >= len(doc.tables):
            return False

        table = doc.tables[table_idx]
        if row_idx >= len(table.rows):
            return False

        row = table.rows[row_idx]
        if cell_idx >= len(row.cells):
            return False

        cell = row.cells[cell_idx]

        # 清除单元格所有内容，然后写入新值
        # 保留第一个段落的格式，删除其他段落
        if cell.paragraphs:
            # 保存第一个段落的格式
            first_para = cell.paragraphs[0]
            first_run = first_para.runs[0] if first_para.runs else None

            # 清除所有段落的文本
            for para in cell.paragraphs:
                for run in para.runs:
                    run.text = ""

            # 删除多余的段落（保留第一个）
            # 注意：python-docx不能直接删除段落，所以我们清空它们
            for para in cell.paragraphs[1:]:
                for run in para.runs:
                    run.text = ""

            # 在第一个段落写入新值
            if first_run:
                first_run.text = value
            elif first_para.runs:
                first_para.runs[0].text = value
            else:
                first_para.text = value
        else:
            cell.text = value

        return True

    def _fill_paragraph_field(self, doc, location: Dict, value: str, fill_mode: str) -> bool:
        """填写段落字段（LLM检测到的）"""
        para_idx = location.get('paragraph_index')
        if para_idx is None or para_idx >= len(doc.paragraphs):
            return False

        para = doc.paragraphs[para_idx]
        text = para.text

        # 检测段落中的占位符模式并填写
        # 1. 检测下划线
        underline_match = re.search(r'_{3,}', text)
        if underline_match:
            return self._fill_underline(doc, location, value, fill_mode)

        # 2. 检测方括号
        bracket_patterns = [
            (re.compile(r'【\s*】'), '【', '】'),
            (re.compile(r'\[\s*\]'), '[', ']'),
            (re.compile(r'（\s*）'), '（', '）'),
        ]
        for pattern, open_bracket, close_bracket in bracket_patterns:
            match = pattern.search(text)
            if match:
                new_text = text[:match.start()] + open_bracket + value + close_bracket + text[match.end():]
                if para.runs:
                    para.runs[0].text = new_text
                    for run in para.runs[1:]:
                        run.text = ""
                else:
                    para.text = new_text
                return True

        # 3. 检测冒号模式（冒号后无内容或只有空格）
        colon_match = re.search(r'[：:]\s*$', text)
        if colon_match:
            return self._fill_colon(doc, location, value, fill_mode)

        # 4. 如果没有检测到占位符，在段落末尾追加值
        if fill_mode == 'append':
            new_text = text + value
        else:
            new_text = text.rstrip() + value

        if para.runs:
            para.runs[0].text = new_text
            for run in para.runs[1:]:
                run.text = ""
        else:
            para.text = new_text

        return True

    def _fill_table_cell_list(self, doc, location: Dict, values: List[Dict[str, str]]) -> bool:
        """填写表格单元格列表（用于多行数据，如成员列表）"""
        table_idx = location.get('table_index')
        header_row_idx = location.get('header_row')

        if table_idx is None or header_row_idx is None:
            return False

        if table_idx >= len(doc.tables):
            return False

        table = doc.tables[table_idx]
        if header_row_idx >= len(table.rows):
            return False

        # 获取表头行
        header_row = table.rows[header_row_idx]
        headers = [cell.text.strip() for cell in header_row.cells]

        # 从表头行下面开始填写
        start_row_idx = header_row_idx + 1

        # 填写数据
        for data_idx, data_row in enumerate(values):
            row_idx = start_row_idx + data_idx

            # 如果行不够，添加新行
            if row_idx >= len(table.rows):
                row = table.add_row()
            else:
                row = table.rows[row_idx]

            # 填写每个单元格
            for col_idx, header in enumerate(headers):
                if col_idx >= len(row.cells):
                    continue

                value = data_row.get(header, "")
                if not value:
                    # 尝试模糊匹配
                    for key, val in data_row.items():
                        if key in header or header in key:
                            value = val
                            break

                if value:
                    cell = row.cells[col_idx]
                    if cell.paragraphs:
                        para = cell.paragraphs[0]
                        if para.runs:
                            para.runs[0].text = str(value)
                            for run in para.runs[1:]:
                                run.text = ""
                        else:
                            para.text = str(value)
                    else:
                        cell.text = str(value)

        return True

    def _fill_content_control(self, doc, location: Dict, value: str) -> bool:
        """填写Word内容控件"""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        sdt_idx = location.get('sdt_index')
        if sdt_idx is None:
            return False

        sdt_elements = doc.element.body.findall('.//' + qn('w:sdt'))
        if sdt_idx >= len(sdt_elements):
            return False

        sdt = sdt_elements[sdt_idx]
        sdt_content = sdt.find(qn('w:sdtContent'))
        if sdt_content is None:
            return False

        # 更新内容控件中的文本
        t_elements = sdt_content.findall('.//' + qn('w:t'))
        if t_elements:
            # 保留第一个文本元素，清空其他
            t_elements[0].text = value
            for t_elem in t_elements[1:]:
                t_elem.text = ""
        else:
            # 没有文本元素，创建 p/r/t 结构（OxmlElement 创建元素，qn 只是名称字符串）
            p_elem = sdt_content.find(qn('w:p'))
            if p_elem is None:
                p_elem = OxmlElement('w:p')
                sdt_content.append(p_elem)
            r_elem = p_elem.find(qn('w:r'))
            if r_elem is None:
                r_elem = OxmlElement('w:r')
                p_elem.append(r_elem)
            new_t = OxmlElement('w:t')
            new_t.text = value
            r_elem.append(new_t)

        return True
