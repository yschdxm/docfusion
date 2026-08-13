"""
文档编辑工具集

将每个编辑操作封装为独立的BaseTool。
采用"复制原始文件 → 在副本上原地修改 → 保存"的方式，保留原始排版格式。
支持：文本替换、段落重写、段落插入、标题调整、列表格式化、段落拆分、样式设置、格式转换。
"""

import csv
import logging
import os
import re
from typing import Any, Dict, List, Optional
from uuid import UUID

from sqlalchemy import select

from app.agent.base.tool import BaseTool, ToolContext, ToolResult
from app.db.postgres import async_session
from app.models.document import Document
from app.services import document_versioning, file_storage
from app.services.document_processor import DocxParser, MdParser, TxtParser, XlsxParser
from app.services.llm_service import llm_service


logger = logging.getLogger(__name__)

CONVERT_RULES = {
    "docx": {"md", "txt"},
    "md": {"docx", "txt"},
    "txt": {"md", "docx"},
    "xlsx": {"csv"},
}


# ──────────────────────────── 公共辅助函数 ────────────────────────────

async def _get_doc_info(file_id: str) -> Optional[Dict[str, Any]]:
    """通过数据库查询文档信息，返回 {file_path, file_type, original_filename, doc_id, is_output}"""
    try:
        doc_uuid = UUID(file_id) if isinstance(file_id, str) else file_id
    except ValueError:
        return None

    async with async_session() as db:
        result = await db.execute(
            select(Document).where(Document.id == doc_uuid)
        )
        doc = result.scalar_one_or_none()
        if doc and doc.file_path and os.path.exists(doc.file_path):
            return {
                "file_path": doc.file_path,
                "file_type": doc.file_type,
                "original_filename": doc.original_filename or os.path.basename(doc.file_path),
                "doc_id": str(doc.id),
                "is_output": (doc.doc_category == "output"),
            }
    return None


async def _resolve_edit_target(doc_info: Dict[str, Any], output_file_id: Optional[str], context) -> tuple:
    """解析编辑目标（版本化）。

    - output 文档：同 run 原地改；新 run 首次编辑创建 version+1（文件已复制好）
    - source/template 首次编辑：创建新 root 的输出 v1（文件已复制好）

    Returns:
        (target_path, original_filename, target_doc_id, is_new_version)
        编辑直接在 target_path 上进行，目标 Document 行已存在。
    """
    run_id = context.metadata.get("run_id") if context else None
    conversation_id = context.metadata.get("conversation_id") if context else None

    async with async_session() as db:
        doc = None
        # 优先使用 LLM 传入的 output_file_id（链式编辑）
        if output_file_id:
            try:
                result = await db.execute(select(Document).where(Document.id == UUID(str(output_file_id))))
                doc = result.scalar_one_or_none()
            except ValueError:
                doc = None
        if doc is None:
            result = await db.execute(select(Document).where(Document.id == UUID(doc_info["doc_id"])))
            doc = result.scalar_one_or_none()
        if doc is None:
            return doc_info["file_path"], doc_info["original_filename"], doc_info["doc_id"], False

        if doc.doc_category == "output":
            target, is_new = await document_versioning.resolve_output_target(
                db, doc,
                run_id=run_id,
                origin_type="edit",
                conversation_id=conversation_id,
            )
        else:
            target = await document_versioning.create_root_output(
                db,
                user_id=doc.user_id,
                file_type=doc.file_type,
                origin_type="edit",
                source_doc=doc,
                run_id=run_id,
                conversation_id=conversation_id,
            )
            is_new = True
        await db.commit()
        return target.file_path, target.original_filename, str(target.id), is_new


# ──────────────────────────── 文本处理辅助 ────────────────────────────

def _find_paragraph_by_index(doc, index: int):
    """获取 doc.paragraphs 中第 index 个非空段落。"""
    editable = [p for p in doc.paragraphs if p.text.strip()]
    if 0 <= index < len(editable):
        return editable[index]
    return None


def _get_editable_paragraphs(doc) -> list:
    """获取所有非空段落列表。"""
    return [p for p in doc.paragraphs if p.text.strip()]


def _insert_paragraph_after(paragraph, text: str = ""):
    """在指定段落后插入新段落并返回（python-docx 无原生 insert_paragraph_after）"""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    new_p = paragraph._p.makeelement(qn("w:p"), {})
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def _set_paragraph_text(paragraph, new_text: str):
    """设置段落文本，保留第一个 run 的格式。"""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = new_text


def _replace_in_paragraph_runs(paragraph, old: str, new: str) -> int:
    """在段落内做 run 感知的文本替换，尽量保留段内混排格式。

    只改写被替换文本覆盖到的 run，其余 run 原样保留；
    新文本继承匹配起点所在 run 的格式。

    Returns:
        替换次数
    """
    if not old or not paragraph.runs:
        return 0

    runs = paragraph.runs
    count = 0
    while True:
        full = "".join(r.text or "" for r in runs)
        idx = full.find(old)
        if idx < 0:
            break
        end = idx + len(old)

        # 定位匹配区间跨越的 run
        pos = 0
        start_run = end_run = None
        start_off = end_off = 0
        for i, r in enumerate(runs):
            rlen = len(r.text or "")
            if start_run is None and idx < pos + rlen:
                start_run, start_off = i, idx - pos
            if start_run is not None and end <= pos + rlen:
                end_run, end_off = i, end - pos
                break
            pos += rlen

        if start_run is None or end_run is None:
            break

        if start_run == end_run:
            r = runs[start_run]
            text = r.text or ""
            r.text = text[:start_off] + new + text[end_off:]
        else:
            first = runs[start_run]
            last = runs[end_run]
            first.text = (first.text or "")[:start_off] + new
            for r in runs[start_run + 1:end_run]:
                r.text = ""
            last.text = (last.text or "")[end_off:]

        count += 1

    return count


def _make_change_record(op_name: str, index: int, before: str, after: str, reason: str) -> Dict[str, Any]:
    return {
        "op": op_name,
        "paragraph_index": index,
        "before": before[:300],
        "after": after[:300],
        "reason": reason,
    }


def _strip_heading_prefix(text: str) -> str:
    return re.sub(r"^#{1,6}\s+", "", text).strip()


def _strip_list_prefix(text: str) -> str:
    return re.sub(r"^(?:[-*]|\d+\.)\s+", "", text).strip()


def _split_paragraph(text: str, separator: Optional[str]) -> List[str]:
    if separator:
        parts = [part.strip() for part in text.split(separator) if part.strip()]
        return parts or [text]
    parts = [part.strip() for part in re.split(r"(?<=[。！？；;.!?])", text) if part.strip()]
    return parts or [text]


def _get_parser(file_type: str):
    """获取文档解析器"""
    parsers = {
        "docx": DocxParser(),
        "xlsx": XlsxParser(),
        "md": MdParser(),
        "txt": TxtParser(),
    }
    return parsers.get(file_type)


# ──────────────────────────── 文档编辑工具 ────────────────────────────

def _add_docx_content(doc, content: str) -> None:
    """按常见 Markdown/纯文本结构写入 docx。"""
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            doc.add_paragraph()
            continue
        heading_match = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading_match:
            doc.add_heading(heading_match.group(2).strip(), level=min(len(heading_match.group(1)), 6))
            continue
        bullet_match = re.match(r"^[-*]\s+(.+)$", line)
        if bullet_match:
            doc.add_paragraph(bullet_match.group(1).strip(), style="List Bullet")
            continue
        number_match = re.match(r"^\d+[.)]\s+(.+)$", line)
        if number_match:
            doc.add_paragraph(number_match.group(1).strip(), style="List Number")
            continue
        doc.add_paragraph(line)


class CreateWordDocumentTool(BaseTool):
    """从零创建 Word 文档工具"""

    @property
    def name(self) -> str:
        return "create_word_document"

    @property
    def description(self) -> str:
        return "未选择模板或没有现成文档时，根据用户提供的标题和正文内容新建一个Word(.docx)文档，并自动保存到文档管理的输出文件中。"

    @property
    def parameters(self) -> Dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "文档标题"},
                "content": {"type": "string", "description": "文档正文，支持简单Markdown标题、列表"},
                "filename": {"type": "string", "description": "输出文件名，可选"},
            },
            "required": ["title", "content"],
        }

    async def execute(self, params: Dict[str, Any], context: ToolContext) -> ToolResult:
        from io import BytesIO

        from docx import Document as DocxDocument
        from docx.oxml.ns import qn
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        title = str(params.get("title", "")).strip()
        content = str(params.get("content", "")).strip()
        if not title:
            return ToolResult(success=False, error="文档标题不能为空")
        if not content:
            return ToolResult(success=False, error="文档正文不能为空")

        requested_filename = params.get("filename") or f"{title}.docx"
        output_filename = file_storage.sanitize_filename(str(requested_filename), ".docx")
        if not output_filename.lower().endswith(".docx"):
            output_filename = f"{output_filename}.docx"

        doc = DocxDocument()
        styles = doc.styles
        styles["Normal"].font.name = "宋体"
        styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        styles["Normal"].font.size = Pt(11)
        title_para = doc.add_heading(title, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_docx_content(doc, content)
        buffer = BytesIO()
        doc.save(buffer)

        user_id = None
        if getattr(context, "user_id", None):
            try:
                user_id = UUID(str(context.user_id))
            except ValueError:
                user_id = None
        async with async_session() as db:
            output_doc = await document_versioning.create_root_output(
                db,
                user_id=user_id,
                file_type="docx",
                origin_type="generate",
                original_filename=output_filename,
                content_bytes=buffer.getvalue(),
                run_id=context.metadata.get("run_id"),
                conversation_id=context.metadata.get("conversation_id"),
            )
            await db.commit()
            reg = document_versioning.download_info(output_doc)
            output_path = output_doc.file_path
        return ToolResult(success=True, data={"message": "已新建Word文档并保存到文档管理", "output_file": output_path, **reg, "output_format": "docx"})


class EditParagraphTool(BaseTool):
    """段落内容编辑工具（合并 replace/rewrite/insert_after/split）"""

    @property
    def name(self) -> str:
        return "edit_paragraph"

    @property
    def description(self) -> str:
        return """编辑文档中指定段落的内容。编辑前必须先用 get_document_outline 获取段落索引[Pn]（禁止凭猜测使用 paragraph_index）。

操作类型（op 参数）：
- replace: 替换段落内文本。需要 old_text + new_text（run感知替换，保留混排格式；old_text 为空时整段替换为 new_text）
- rewrite: 用LLM重写整段。需要 rewrite_instruction
- insert_after: 在该段后插入新段落。需要 text（可选 style 指定段落样式）
- split: 将长段拆分为多段。可选 separator（默认按句号拆分）

通用参数：
- file_id: 文档ID（首次用原始ID，后续用上次返回的 output_file_id）
- paragraph_index: 目标段落索引（从0开始，与 get_document_outline 的 [Pn] 一致；md/txt 为行号[Ln]）
- 支持 docx / md / txt；表格单元格内容请用 edit_docx_cell 或 find_replace_all"""

    @property
    def parameters(self) -> Dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "file_id": {"type": "string", "description": "文档ID（首次用原始ID，后续用上次返回的output_file_id）"},
                "output_file_id": {"type": "string", "description": "链式编辑时上次返回的 output_file_id（可选）"},
                "op": {"type": "string", "enum": ["replace", "rewrite", "insert_after", "split"], "description": "操作类型"},
                "paragraph_index": {"type": "integer", "description": "目标段落索引（从0开始，来自 get_document_outline 的 [Pn]）"},
                "old_text": {"type": "string", "description": "（op=replace）要替换的原文本"},
                "new_text": {"type": "string", "description": "（op=replace）替换后的新文本"},
                "rewrite_instruction": {"type": "string", "description": "（op=rewrite）重写指令"},
                "text": {"type": "string", "description": "（op=insert_after）要插入的文本"},
                "style": {"type": "string", "description": "（op=insert_after）段落样式（可选，默认Normal）"},
                "separator": {"type": "string", "description": "（op=split）分隔符（可选，默认按句号拆分）"},
                "reason": {"type": "string", "description": "编辑原因（可选）"},
            },
            "required": ["file_id", "op", "paragraph_index"],
        }

    async def execute(self, params: Dict[str, Any], context: ToolContext) -> ToolResult:
        doc_info = await _get_doc_info(params["file_id"])
        if not doc_info:
            return ToolResult(success=False, error=f"文档 {params['file_id']} 不存在")

        op = params.get("op", "")
        paragraph_index = params["paragraph_index"]
        output_file_id = params.get("output_file_id")

        target_path, _, target_doc_id, _ = await _resolve_edit_target(doc_info, output_file_id, context)

        if doc_info["file_type"] == "docx":
            result = await self._apply_docx(target_path, op, paragraph_index, params)
        else:
            result = self._apply_text(target_path, op, paragraph_index, params)

        if not result["ok"]:
            return ToolResult(success=False, error=result["error"])
        if result.get("no_change"):
            return ToolResult(success=True, data={"message": result["message"]})

        reg = await _finish_edit(target_doc_id, context)
        return ToolResult(
            success=True,
            data={
                "message": result["message"],
                "output_file": target_path,
                **reg,
                "changes": [_make_change_record(f"edit_paragraph:{op}", paragraph_index, result.get("before", ""), result.get("after", ""), params.get("reason", ""))],
            },
        )

    async def _apply_docx(self, target_path: str, op: str, paragraph_index: int, params: Dict[str, Any]) -> Dict[str, Any]:
        from docx import Document as DocxDocument

        doc = DocxDocument(target_path)
        editable = _get_editable_paragraphs(doc)
        if not (0 <= paragraph_index < len(editable)):
            return {"ok": False, "error": f"段落索引 {paragraph_index} 超出范围（共 {len(editable)} 个非空段落）"}
        paragraph = editable[paragraph_index]
        before = paragraph.text

        if op == "replace":
            old_text = params.get("old_text", "")
            new_text = params.get("new_text", "")
            if old_text:
                replaced = _replace_in_paragraph_runs(paragraph, old_text, new_text)
                if replaced == 0:
                    return {"ok": False, "error": f"在第 {paragraph_index} 段中未找到文本: '{old_text[:50]}'。当前段落内容: '{before[:100]}'"}
            else:
                _set_paragraph_text(paragraph, new_text)
            after = paragraph.text
            message = f"已在第 {paragraph_index} 段中替换文本"

        elif op == "rewrite":
            instruction = params.get("rewrite_instruction", "请重写此段落。")
            after = await llm_service.rewrite_paragraph_text(before, instruction)
            _set_paragraph_text(paragraph, after)
            message = f"已重写第 {paragraph_index} 段"

        elif op == "insert_after":
            insert_text = params.get("text", "").strip()
            if not insert_text:
                return {"ok": False, "error": "插入文本不能为空（text 参数）"}
            new_para = _insert_paragraph_after(paragraph, insert_text)
            style = params.get("style", "Normal")
            if style and style != "Normal":
                try:
                    new_para.style = style
                except (KeyError, ValueError):
                    pass
            before, after = "", insert_text
            message = f"已在第 {paragraph_index} 段后插入新内容"

        elif op == "split":
            parts = _split_paragraph(before, params.get("separator"))
            if len(parts) <= 1:
                return {"ok": True, "no_change": True, "message": "段落无法拆分（未找到分隔点）"}
            style = paragraph.style
            _set_paragraph_text(paragraph, parts[0])
            current_para = paragraph
            for part in parts[1:]:
                new_para = _insert_paragraph_after(current_para, part)
                try:
                    new_para.style = style
                except (KeyError, ValueError):
                    pass
                if paragraph.runs:
                    ref_run = paragraph.runs[0]
                    for run in new_para.runs:
                        run.font.name = ref_run.font.name
                        run.font.size = ref_run.font.size
                        run.font.bold = ref_run.font.bold
                        run.font.italic = ref_run.font.italic
                current_para = new_para
            after = "\n".join(parts)
            message = f"已将第 {paragraph_index} 段拆分为 {len(parts)} 个段落"

        else:
            return {"ok": False, "error": f"不支持的操作类型: {op}（可选 replace/rewrite/insert_after/split）"}

        doc.save(target_path)
        return {"ok": True, "message": message, "before": before, "after": after}

    def _apply_text(self, target_path: str, op: str, paragraph_index: int, params: Dict[str, Any]) -> Dict[str, Any]:
        with open(target_path, "r", encoding="utf-8") as f:
            lines = [line.rstrip("\n") for line in f.readlines()]

        if not (0 <= paragraph_index < len(lines)):
            return {"ok": False, "error": f"段落索引 {paragraph_index} 超出范围（共 {len(lines)} 行）"}
        before = lines[paragraph_index]

        if op == "replace":
            old_text = params.get("old_text", "")
            new_text = params.get("new_text", "")
            if old_text and old_text not in before:
                return {"ok": False, "error": f"在第 {paragraph_index} 行中未找到文本: '{old_text[:50]}'。当前行内容: '{before[:100]}'"}
            lines[paragraph_index] = before.replace(old_text, new_text) if old_text else new_text
            after = lines[paragraph_index]
            message = f"已在第 {paragraph_index} 行中替换文本"

        elif op == "rewrite":
            return {"ok": False, "error": "rewrite 操作仅支持 docx 文件，md/txt 请用 replace"}

        elif op == "insert_after":
            insert_text = params.get("text", "").strip()
            if not insert_text:
                return {"ok": False, "error": "插入文本不能为空（text 参数）"}
            lines.insert(paragraph_index + 1, insert_text)
            before, after = "", insert_text
            message = f"已在第 {paragraph_index} 行后插入新内容"

        elif op == "split":
            parts = _split_paragraph(before, params.get("separator"))
            if len(parts) <= 1:
                return {"ok": True, "no_change": True, "message": "段落无法拆分（未找到分隔点）"}
            lines[paragraph_index] = parts[0]
            for offset, part in enumerate(parts[1:], start=1):
                lines.insert(paragraph_index + offset, part)
            after = "\n".join(parts)
            message = f"已将第 {paragraph_index} 行拆分为 {len(parts)} 行"

        else:
            return {"ok": False, "error": f"不支持的操作类型: {op}（可选 replace/rewrite/insert_after/split）"}

        with open(target_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        return {"ok": True, "message": message, "before": before, "after": after}


class FormatParagraphTool(BaseTool):
    """段落格式工具（合并 heading/list/style）"""

    @property
    def name(self) -> str:
        return "format_paragraph"

    @property
    def description(self) -> str:
        return """调整指定段落的格式。编辑前必须先用 get_document_outline 获取段落索引[Pn]。

操作类型（op 参数）：
- heading: 设为标题。需要 level（1-6）。docx 设为 Heading N 样式；md/txt 加 # 前缀
- list: 转为列表。需要 list_type（bullet/number），按 paragraph_indexes 顺序编号
- style: 设置字体样式（仅 docx）。可选 font_name、font_size_pt

通用参数：
- file_id: 文档ID（首次用原始ID，后续用上次返回的 output_file_id）
- paragraph_indexes: 段落索引数组（从0开始，来自 get_document_outline 的 [Pn]）"""

    @property
    def parameters(self) -> Dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "file_id": {"type": "string", "description": "文档ID（首次用原始ID，后续用上次返回的output_file_id）"},
                "output_file_id": {"type": "string", "description": "链式编辑时上次返回的 output_file_id（可选）"},
                "op": {"type": "string", "enum": ["heading", "list", "style"], "description": "操作类型"},
                "paragraph_indexes": {"type": "array", "items": {"type": "integer"}, "description": "段落索引数组（来自 get_document_outline 的 [Pn]）"},
                "level": {"type": "integer", "description": "（op=heading）标题级别 1-6"},
                "list_type": {"type": "string", "enum": ["bullet", "number"], "description": "（op=list）列表类型"},
                "font_name": {"type": "string", "description": "（op=style）字体名称（可选）"},
                "font_size_pt": {"type": "number", "description": "（op=style）字体大小 pt（可选）"},
                "reason": {"type": "string", "description": "操作原因（可选）"},
            },
            "required": ["file_id", "op", "paragraph_indexes"],
        }

    async def execute(self, params: Dict[str, Any], context: ToolContext) -> ToolResult:
        doc_info = await _get_doc_info(params["file_id"])
        if not doc_info:
            return ToolResult(success=False, error=f"文档 {params['file_id']} 不存在")

        op = params.get("op", "")
        indexes = params.get("paragraph_indexes", [])
        if not indexes:
            return ToolResult(success=False, error="paragraph_indexes 不能为空")

        if op == "style" and doc_info["file_type"] != "docx":
            return ToolResult(success=False, error="字体样式设置仅支持docx文件")

        output_file_id = params.get("output_file_id")
        target_path, _, target_doc_id, _ = await _resolve_edit_target(doc_info, output_file_id, context)

        if doc_info["file_type"] == "docx":
            changes, message = self._apply_docx(target_path, op, indexes, params)
        else:
            changes, message = self._apply_text(target_path, op, indexes, params)

        if changes is None:
            return ToolResult(success=False, error=message)

        reg = await _finish_edit(target_doc_id, context)
        return ToolResult(
            success=True,
            data={
                "message": message,
                "output_file": target_path,
                **reg,
                "changes": changes,
            },
        )

    def _apply_docx(self, target_path: str, op: str, indexes: List[int], params: Dict[str, Any]) -> tuple:
        from docx.shared import Pt
        from docx import Document as DocxDocument

        doc = DocxDocument(target_path)
        editable = _get_editable_paragraphs(doc)
        changes = []

        if op == "heading":
            level = max(1, min(6, int(params.get("level", 1))))
            for index in indexes:
                if not (0 <= index < len(editable)):
                    continue
                paragraph = editable[index]
                before = f"[{paragraph.style.name if paragraph.style else 'Normal'}] {paragraph.text}"
                paragraph.style = f"Heading {level}"
                after = f"[{paragraph.style.name}] {paragraph.text}"
                changes.append(_make_change_record("format_paragraph:heading", index, before, after, params.get("reason", "")))
            message = f"已将 {len(changes)} 个段落调整为 Heading {level}"

        elif op == "list":
            list_type = params.get("list_type", "bullet")
            for position, index in enumerate(indexes, start=1):
                if not (0 <= index < len(editable)):
                    continue
                paragraph = editable[index]
                before = paragraph.text
                clean_text = _strip_list_prefix(before)
                if list_type == "number":
                    _set_paragraph_text(paragraph, f"{position}. {clean_text}")
                    try:
                        paragraph.style = "List Number"
                    except (KeyError, ValueError):
                        pass
                else:
                    _set_paragraph_text(paragraph, f"- {clean_text}")
                    try:
                        paragraph.style = "List Bullet"
                    except (KeyError, ValueError):
                        pass
                changes.append(_make_change_record("format_paragraph:list", index, before, paragraph.text, params.get("reason", "")))
            message = f"已将 {len(changes)} 个段落格式化为{'有序' if list_type == 'number' else '无序'}列表"

        elif op == "style":
            font_name = params.get("font_name", "")
            font_size_pt = params.get("font_size_pt")
            for index in indexes:
                if not (0 <= index < len(editable)):
                    continue
                paragraph = editable[index]
                before = f"[{paragraph.style.name if paragraph.style else 'Normal'}] {paragraph.text}"
                for run in paragraph.runs:
                    if font_name:
                        run.font.name = font_name
                    if font_size_pt is not None:
                        run.font.size = Pt(font_size_pt)
                after = f"[{paragraph.style.name if paragraph.style else 'Normal'}] {paragraph.text}"
                changes.append(_make_change_record("format_paragraph:style", index, before, after, params.get("reason", "")))
            message = f"已设置 {len(changes)} 个段落的字体样式"

        else:
            return None, f"不支持的操作类型: {op}（可选 heading/list/style）"

        doc.save(target_path)
        return changes, message

    def _apply_text(self, target_path: str, op: str, indexes: List[int], params: Dict[str, Any]) -> tuple:
        with open(target_path, "r", encoding="utf-8") as f:
            lines = [line.rstrip("\n") for line in f.readlines()]
        changes = []

        if op == "heading":
            level = max(1, min(6, int(params.get("level", 1))))
            for index in indexes:
                if not (0 <= index < len(lines)):
                    continue
                before = lines[index]
                clean = _strip_heading_prefix(before)
                lines[index] = f"{'#' * level} {clean}"
                changes.append(_make_change_record("format_paragraph:heading", index, before, lines[index], params.get("reason", "")))
            message = f"已将 {len(changes)} 行调整为 {level} 级标题"

        elif op == "list":
            list_type = params.get("list_type", "bullet")
            for position, index in enumerate(indexes, start=1):
                if not (0 <= index < len(lines)):
                    continue
                before = lines[index]
                clean_text = _strip_list_prefix(before)
                lines[index] = f"{position}. {clean_text}" if list_type == "number" else f"- {clean_text}"
                changes.append(_make_change_record("format_paragraph:list", index, before, lines[index], params.get("reason", "")))
            message = f"已将 {len(changes)} 行格式化为{'有序' if list_type == 'number' else '无序'}列表"

        else:
            return None, f"不支持的操作类型: {op}（md/txt 支持 heading/list；style 仅 docx）"

        with open(target_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        return changes, message



class ConvertTool(BaseTool):
    """格式转换工具"""

    @property
    def name(self) -> str:
        return "convert"

    @property
    def description(self) -> str:
        return "转换文档格式。支持：docx<->md, docx<->txt, md<->txt, xlsx->csv。"

    @property
    def parameters(self) -> Dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "file_id": {"type": "string", "description": "文档ID"},
                "target_format": {"type": "string", "enum": ["docx", "md", "txt", "csv"], "description": "目标格式"},
                "reason": {"type": "string", "description": "转换原因（可选）"},
            },
            "required": ["file_id", "target_format"],
        }

    async def execute(self, params: Dict[str, Any], context: ToolContext) -> ToolResult:
        doc_info = await _get_doc_info(params["file_id"])
        if not doc_info:
            return ToolResult(success=False, error=f"文档 {params['file_id']} 不存在")

        file_type = doc_info["file_type"]
        target_format = params.get("target_format", "").lower()

        if target_format not in CONVERT_RULES.get(file_type, set()):
            return ToolResult(success=False, error=f"不支持从 {file_type} 转换为 {target_format}")

        parser = _get_parser(file_type)
        if not parser:
            return ToolResult(success=False, error=f"不支持的文件类型: {file_type}")

        parsed_data = parser.parse(doc_info["file_path"])

        # 先创建新 root 的输出行，再把转换内容写入其版本文件
        base_name = doc_info["original_filename"].rsplit(".", 1)[0]
        output_filename = f"{base_name}.{target_format}"
        user_id = None
        if getattr(context, "user_id", None):
            try:
                user_id = UUID(str(context.user_id))
            except ValueError:
                user_id = None
        async with async_session() as db:
            result = await db.execute(
                select(Document).where(Document.id == UUID(doc_info["doc_id"]))
            )
            source_doc = result.scalar_one_or_none()
            output_doc = await document_versioning.create_root_output(
                db,
                user_id=user_id or (source_doc.user_id if source_doc else None),
                file_type=target_format,
                origin_type="convert",
                source_doc=None,  # 内容由转换生成，不复制源文件
                original_filename=output_filename,
                run_id=context.metadata.get("run_id"),
                conversation_id=context.metadata.get("conversation_id"),
                extra_metadata={"derived_from": doc_info["doc_id"]},
            )
            output_path = output_doc.file_path

            if target_format == "docx":
                content = parsed_data.get("full_text", "")
                DocxParser.write(content, output_path, {
                    "paragraphs": parsed_data.get("paragraphs", []),
                    "tables": parsed_data.get("tables", []),
                })
            elif target_format == "md":
                MdParser.write(parsed_data.get("full_text", ""), output_path)
            elif target_format == "txt":
                TxtParser.write(parsed_data.get("full_text", ""), output_path)
            elif target_format == "csv":
                rows = []
                sheets = parsed_data.get("sheets", [])
                if sheets:
                    rows = sheets[0].get("data", [])
                with open(output_path, "w", encoding="utf-8-sig", newline="") as csv_file:
                    writer = csv.writer(csv_file)
                    for row in rows:
                        writer.writerow(row)

            output_doc.file_size = os.path.getsize(output_path)
            output_doc.sha256 = file_storage.sha256_file(output_path)
            await db.commit()
            reg = document_versioning.download_info(output_doc)

        return ToolResult(
            success=True,
            data={
                "message": f"已将文档从 {file_type} 转换为 {target_format}",
                "output_file": output_path,
                **reg,
                "output_format": target_format,
            },
        )


# ──────────────────────────── 工具共用的完成处理 ────────────────────────────

async def _finish_edit(target_doc_id: str, context) -> Dict[str, str]:
    """编辑完成后更新目标版本行的大小/哈希，返回下载信息（键保持不变）。"""
    async with async_session() as db:
        result = await db.execute(
            select(Document).where(Document.id == UUID(str(target_doc_id)))
        )
        doc = result.scalar_one_or_none()
        if not doc:
            return {
                "output_file_id": str(target_doc_id),
                "output_filename": "",
                "download_url": f"/api/v1/documents/{target_doc_id}/download",
            }
        if doc.file_path and os.path.exists(doc.file_path):
            doc.file_size = os.path.getsize(doc.file_path)
            doc.sha256 = file_storage.sha256_file(doc.file_path)
        doc.status = "completed"
        await db.commit()
        return document_versioning.download_info(doc)
