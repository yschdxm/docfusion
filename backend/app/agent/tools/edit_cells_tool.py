"""
单元格级编辑工具集

提供原子写能力：
- edit_xlsx_cells: 按 sheet + 坐标改写 Excel 任意单元格
- edit_docx_cell: 按 table/row/cell 索引改写 Word 任意表格单元格
- find_replace_all: 在 Word/文本文档全域（段落+表格）查找替换，无需索引

版本化约定：编辑目标由 _resolve_edit_target 解析（同 run 原地改，
新 run 自动创建新版本），编辑直接作用于目标版本文件。
首次编辑用原始文档ID，后续编辑用返回的 output_file_id。
"""

import logging
from typing import Any, Dict, List, Optional

from app.agent.base.tool import BaseTool, ToolContext, ToolResult
from app.agent.document_engine import lifecycle, ops
from app.agent.document_engine.address import (
    AddressResolver,
    DocxTableCellAddr,
    XlsxCellAddr,
)
from app.agent.document_engine.snapshot import DocSnapshot

_get_doc_info = lifecycle.get_doc_info
_resolve_edit_target = lifecycle.resolve_edit_target
_finish_edit = lifecycle.finish_edit

logger = logging.getLogger(__name__)


class EditXlsxCellsTool(BaseTool):
    """Excel 单元格编辑工具

    按坐标精确改写任意单元格，openpyxl 原地写入，保留其他单元格和格式。
    """

    @property
    def name(self) -> str:
        return "edit_xlsx_cells"

    @property
    def description(self) -> str:
        return """精确改写Excel文件中的指定单元格（原子写能力）。

使用场景：
- 修改某个具体单元格（如 B5）而不影响其他数据
- 填写纵向表单式Excel模板（A列标签、B列填值），逐个字段写入
- 向指定工作表、指定位置写入数据
- 修正 fill_table 填写后的个别错误单元格

参数说明：
- file_id: 文档ID（首次编辑用原始ID，后续编辑用上一步返回的 output_file_id）
- sheet_name: 工作表名称（可选，默认当前活动工作表）
- edits: 编辑列表，每项支持两种坐标写法：
  - {"cell": "B5", "value": "新值"}
  - {"row": 5, "col": 2, "value": "新值"}  （行列均从1开始）
  - 可选 "value_type": "number" 强制按数值写入（默认保留传入类型）
  - value 传 null 表示清空该单元格
- 编辑前建议先用 read_document(sheet_name=..., cell_range=...) 查看目标区域当前内容"""

    @property
    def parameters(self) -> Dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "file_id": {
                    "type": "string",
                    "description": "文档ID（首次用原始ID，后续用上次返回的output_file_id）"
                },
                "output_file_id": {
                    "type": "string",
                    "description": "链式编辑时上次返回的 output_file_id（可选，与 file_id 二选一）"
                },
                "sheet_name": {
                    "type": "string",
                    "description": "工作表名称（可选，默认活动工作表）"
                },
                "edits": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "cell": {"type": "string", "description": "单元格坐标，如 B5"},
                            "row": {"type": "integer", "description": "行号（从1开始）"},
                            "col": {"type": "integer", "description": "列号（从1开始）"},
                            "sheet": {"type": "string", "description": "该项单独指定工作表（可选，覆盖sheet_name）"},
                            "value": {"description": "新值，null表示清空"},
                            "value_type": {"type": "string", "enum": ["string", "number"], "description": "强制值类型（可选）"}
                        }
                    },
                    "description": "编辑列表"
                },
                "reason": {"type": "string", "description": "编辑原因（可选）"}
            },
            "required": ["file_id", "edits"]
        }

    async def execute(self, params: Dict[str, Any], context: ToolContext) -> ToolResult:
        doc_info = await _get_doc_info(params["file_id"])
        if not doc_info:
            return ToolResult(success=False, error=f"文档 {params['file_id']} 不存在")

        if doc_info["file_type"] != "xlsx":
            return ToolResult(
                success=False,
                error=f"edit_xlsx_cells 仅支持xlsx文件，当前格式: {doc_info['file_type']}"
            )

        edits = params.get("edits", [])
        if not edits:
            return ToolResult(success=False, error="编辑列表不能为空")

        sheet_name = params.get("sheet_name")
        output_file_id = params.get("output_file_id")
        target_path, _, target_doc_id, _ = await _resolve_edit_target(doc_info, output_file_id, context)

        try:
            from openpyxl import load_workbook
            from openpyxl.utils import get_column_letter

            wb = load_workbook(target_path)
            snapshot = DocSnapshot(target_path, "xlsx", wb=wb)

            applied: List[Dict[str, Any]] = []
            errors: List[str] = []

            for i, edit in enumerate(edits):
                # 归一化为引擎地址（row/col 转 A1 坐标）
                coord = edit.get("cell")
                if not coord and edit.get("row") and edit.get("col"):
                    coord = f"{get_column_letter(int(edit['col']))}{int(edit['row'])}"
                if not coord:
                    errors.append(f"第{i + 1}项: 缺少坐标（需要 cell 或 row+col）")
                    continue

                addr = XlsxCellAddr(sheet=edit.get("sheet") or sheet_name, cell=coord)
                resolved = AddressResolver.resolve(addr, snapshot)
                if not resolved.ok:
                    errors.append(f"第{i + 1}项: {resolved.detail}")
                    continue

                result = ops.set_xlsx_cell(resolved.target, edit.get("value"), edit.get("value_type"))
                if result.ok:
                    applied.append({
                        "sheet": addr.sheet or wb.active.title,
                        "cell": resolved.target.coordinate,
                        "before": result.before,
                        "after": result.after,
                    })
                else:
                    errors.append(f"第{i + 1}项: {result.error}")

            if not applied:
                wb.close()
                return ToolResult(
                    success=False,
                    error=f"没有成功应用任何编辑。错误: {'; '.join(errors)}"
                )

            wb.save(target_path)
            wb.close()

        except Exception as e:
            logger.exception(f"编辑Excel单元格失败: {e}")
            return ToolResult(success=False, error=f"编辑Excel单元格失败: {str(e)}")

        reg = await _finish_edit(target_doc_id, context)

        result_data: Dict[str, Any] = {
            "message": f"已修改 {len(applied)} 个单元格",
            "output_file": target_path,
            **reg,
            "changes": applied,
        }
        if errors:
            result_data["partial_errors"] = errors

        return ToolResult(success=True, data=result_data)


class EditDocxCellTool(BaseTool):
    """Word 表格单元格编辑工具

    按 table/row/cell 索引精确改写任意表格单元格。
    """

    @property
    def name(self) -> str:
        return "edit_docx_cell"

    @property
    def description(self) -> str:
        return """精确改写Word文档中指定表格的指定单元格（原子写能力）。

使用场景：
- 修改表格中某个具体单元格的文字
- 修正 fill_table / fill_form 填写后的个别错误单元格

参数说明：
- file_id: 文档ID（首次编辑用原始ID，后续编辑用上一步返回的 output_file_id）
- table_index: 表格索引（从0开始）
- row_index: 行索引（从0开始，第0行通常是表头）
- cell_index: 列索引（从0开始）
- value: 新值（空字符串表示清空）
- mode: replace(覆盖，默认) 或 append(追加到现有内容后面)
- 编辑前建议先用 get_document_outline 获取表格和单元格索引"""

    @property
    def parameters(self) -> Dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "file_id": {"type": "string", "description": "文档ID（首次用原始ID，后续用上次返回的output_file_id）"},
                "output_file_id": {"type": "string", "description": "链式编辑时上次返回的 output_file_id（可选）"},
                "table_index": {"type": "integer", "description": "表格索引（从0开始）"},
                "row_index": {"type": "integer", "description": "行索引（从0开始）"},
                "cell_index": {"type": "integer", "description": "列索引（从0开始）"},
                "value": {"type": "string", "description": "新值"},
                "mode": {"type": "string", "enum": ["replace", "append"], "default": "replace", "description": "写入模式"},
                "reason": {"type": "string", "description": "编辑原因（可选）"}
            },
            "required": ["file_id", "table_index", "row_index", "cell_index", "value"]
        }

    async def execute(self, params: Dict[str, Any], context: ToolContext) -> ToolResult:
        doc_info = await _get_doc_info(params["file_id"])
        if not doc_info:
            return ToolResult(success=False, error=f"文档 {params['file_id']} 不存在")

        if doc_info["file_type"] != "docx":
            return ToolResult(
                success=False,
                error=f"edit_docx_cell 仅支持docx文件，当前格式: {doc_info['file_type']}"
            )

        table_index = params["table_index"]
        row_index = params["row_index"]
        cell_index = params["cell_index"]
        value = params["value"]
        mode = params.get("mode", "replace")
        output_file_id = params.get("output_file_id")

        target_path, _, target_doc_id, _ = await _resolve_edit_target(doc_info, output_file_id, context)

        from docx import Document as DocxDocument
        doc = DocxDocument(target_path)

        # 引擎寻址 + 统一单元格写入（保留首 run 格式）
        snapshot = DocSnapshot(target_path, "docx", docx=doc)
        addr = DocxTableCellAddr(table_index=table_index, row=row_index, col=cell_index)
        resolved = AddressResolver.resolve(addr, snapshot)
        if not resolved.ok:
            return ToolResult(success=False, error=resolved.detail)

        result = ops.set_cell_text(resolved.target, value, mode=mode)
        before = result.before
        new_text = result.after

        doc.save(target_path)

        reg = await _finish_edit(target_doc_id, context)

        return ToolResult(
            success=True,
            data={
                "message": f"已将 [T{table_index}R{row_index}C{cell_index}] 从 '{before[:50]}' 修改为 '{new_text[:50]}'",
                "output_file": target_path,
                **reg,
                "changes": [{
                    "op": "edit_docx_cell",
                    "table_index": table_index,
                    "row_index": row_index,
                    "cell_index": cell_index,
                    "before": before[:300],
                    "after": new_text[:300],
                    "reason": params.get("reason", ""),
                }],
            },
        )


class FindReplaceAllTool(BaseTool):
    """全文查找替换工具

    在文档全域（正文段落 + 表格单元格）查找并替换文本，无需索引。
    """

    @property
    def name(self) -> str:
        return "find_replace_all"

    @property
    def description(self) -> str:
        return """在整个文档中查找并替换指定文本（正文段落和所有表格单元格都会搜索）。

使用场景：
- 不知道目标文本在哪个段落/单元格时的全局替换
- 批量统一术语、姓名、日期等
- 替换表格内的文字（replace_text 类工具无法触及表格内容）

参数说明：
- file_id: 文档ID（首次编辑用原始ID，后续编辑用上一步返回的 output_file_id）
- old_text: 要查找的原文本（必须精确匹配）
- new_text: 替换后的新文本
- max_replacements: 最多替换次数（可选，默认替换全部）
- 支持 docx / md / txt；xlsx 请用 edit_xlsx_cells 按坐标修改

注意：替换是全局的，old_text 出现多处时会全部替换，请确认 old_text 足够具体"""

    @property
    def parameters(self) -> Dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "file_id": {"type": "string", "description": "文档ID（首次用原始ID，后续用上次返回的output_file_id）"},
                "output_file_id": {"type": "string", "description": "链式编辑时上次返回的 output_file_id（可选）"},
                "old_text": {"type": "string", "description": "要查找的原文本"},
                "new_text": {"type": "string", "description": "替换后的新文本"},
                "max_replacements": {"type": "integer", "description": "最多替换次数（可选，默认全部替换）"},
                "reason": {"type": "string", "description": "替换原因（可选）"}
            },
            "required": ["file_id", "old_text", "new_text"]
        }

    async def execute(self, params: Dict[str, Any], context: ToolContext) -> ToolResult:
        doc_info = await _get_doc_info(params["file_id"])
        if not doc_info:
            return ToolResult(success=False, error=f"文档 {params['file_id']} 不存在")

        old_text = params["old_text"]
        new_text = params["new_text"]
        max_replacements: Optional[int] = params.get("max_replacements")

        if not old_text:
            return ToolResult(success=False, error="old_text 不能为空")

        file_type = doc_info["file_type"]
        output_file_id = params.get("output_file_id")
        target_path, _, target_doc_id, _ = await _resolve_edit_target(doc_info, output_file_id, context)

        if file_type == "docx":
            return await self._replace_in_docx(
                target_path, target_doc_id,
                old_text, new_text, max_replacements, context
            )
        elif file_type in ("md", "txt"):
            return await self._replace_in_text(
                target_path, target_doc_id, file_type,
                old_text, new_text, max_replacements, context
            )
        else:
            return ToolResult(
                success=False,
                error=f"find_replace_all 支持 docx/md/txt，当前格式: {file_type}。xlsx 请使用 edit_xlsx_cells"
            )

    async def _replace_in_docx(
        self, target_path, target_doc_id,
        old_text, new_text, max_replacements, context
    ) -> ToolResult:
        from docx import Document as DocxDocument

        doc = DocxDocument(target_path)

        total_replaced = 0
        locations: List[str] = []
        limit = max_replacements if max_replacements else float("inf")

        def _replace_in_paragraphs(paragraphs, location_prefix):
            nonlocal total_replaced
            for para in paragraphs:
                if total_replaced >= limit:
                    return
                if old_text in para.text:
                    n = ops.replace_in_paragraph_runs(para, old_text, new_text)
                    if n > 0:
                        locations.append(location_prefix)
                        total_replaced += n

        # 正文段落
        _replace_in_paragraphs(doc.paragraphs, "正文段落")

        # 表格单元格（按元素去重，合并单元格会重复出现）
        for t_idx, table in enumerate(doc.tables):
            seen_cells = set()
            for r_idx, row in enumerate(table.rows):
                for c_idx, cell in enumerate(row.cells):
                    elem_id = id(cell._element)
                    if elem_id in seen_cells:
                        continue
                    seen_cells.add(elem_id)
                    _replace_in_paragraphs(cell.paragraphs, f"[T{t_idx}R{r_idx}C{c_idx}]")

        if total_replaced == 0:
            return ToolResult(
                success=False,
                error=f"在文档中未找到文本: '{old_text[:80]}'"
            )

        doc.save(target_path)
        reg = await _finish_edit(target_doc_id, context)

        return ToolResult(
            success=True,
            data={
                "message": f"已将 '{old_text[:50]}' 替换为 '{new_text[:50]}'，共 {total_replaced} 处",
                "output_file": target_path,
                **reg,
                "replacements": total_replaced,
                "locations": locations[:20],
            },
        )

    async def _replace_in_text(
        self, target_path, target_doc_id, file_type,
        old_text, new_text, max_replacements, context
    ) -> ToolResult:
        with open(target_path, "r", encoding="utf-8") as f:
            content = f.read()

        occurrences = content.count(old_text)
        if occurrences == 0:
            return ToolResult(
                success=False,
                error=f"在文档中未找到文本: '{old_text[:80]}'"
            )

        if max_replacements:
            content = content.replace(old_text, new_text, max_replacements)
            replaced = min(occurrences, max_replacements)
        else:
            content = content.replace(old_text, new_text)
            replaced = occurrences

        with open(target_path, "w", encoding="utf-8") as f:
            f.write(content)

        reg = await _finish_edit(target_doc_id, context)

        return ToolResult(
            success=True,
            data={
                "message": f"已将 '{old_text[:50]}' 替换为 '{new_text[:50]}'，共 {replaced} 处",
                "output_file": target_path,
                **reg,
                "replacements": replaced,
            },
        )
