"""
模板结构分析工具 - 填表/填表单一站式结构探测

取代旧的 get_template_type / get_table_structure / get_form_structure 三个工具：
- 启发式判断模板类型（data_table / form / mixed / unknown），无 LLM 调用
- xlsx: 自动探测真实表头行（处理标题行/说明行）、合并单元格、多工作表；
  识别纵向表单（A列标签 B列填值）并给出每个字段的写入坐标
- docx: 表格结构（表头、示例、表格上下文）；表单字段检测
  （启发式占位符扫描 + 长文档分块 LLM 分析，字段带稳定 field_id）

探测结果写入 structure_cache（键 = doc_id + sha256），
fill_form 直接复用同一份字段列表，不再重复 LLM 检测。
"""

import json
import re
from typing import Any, Dict, List, Optional, Tuple

from sqlalchemy import select

from app.agent.base.tool import BaseTool, ToolContext, ToolResult
from app.agent.tools import structure_cache
from app.db.postgres import async_session
from app.models.document import Document


# ──────────────────────────── 通用启发式 ────────────────────────────

def detect_xlsx_header_row(rows: List[Tuple], max_scan: int = 10) -> int:
    """探测 Excel 的真实表头行（处理标题行、说明行等情况）。

    打分规则：
    - 非空率高的行更可能是表头
    - 文本比例高的行更可能是表头（数据行通常含数值）
    - 下一行非空数与当前行接近的，更可能是"表头→数据"的衔接

    Args:
        rows: 从第1行开始的行数据列表（values_only）
        max_scan: 最多扫描的行数

    Returns:
        表头行号（1-based）
    """
    best_row, best_score = 1, -1.0

    for i, row in enumerate(rows[:max_scan], start=1):
        cells = list(row)
        non_empty = [v for v in cells if v is not None and str(v).strip() != ""]
        if not non_empty:
            continue

        fill_rate = len(non_empty) / max(len(cells), 1)
        text_rate = sum(1 for v in non_empty if isinstance(v, str)) / len(non_empty)

        # 下一行是否有数据（表头的下一行通常有数据，且列覆盖接近）
        next_bonus = 0.0
        if i < len(rows):
            next_non_empty = sum(
                1 for v in rows[i] if v is not None and str(v).strip() != ""
            )
            if next_non_empty >= len(non_empty) * 0.5:
                next_bonus = 1.0

        score = fill_rate * 2 + text_rate + next_bonus
        if score > best_score:
            best_score, best_row = score, i

    return best_row


def _infer_field_type(label: str) -> str:
    """根据标签文本推断字段类型"""
    type_keywords = {
        'date': ['日期', '时间', '年', '月', '日', '出生', 'date', 'time'],
        'phone': ['电话', '手机', '联系方式', 'phone', 'tel'],
        'email': ['邮箱', '邮件', 'email'],
        'number': ['编号', '号码', '序号', '数量', '金额', '价格', 'number', 'id'],
        'address': ['地址', '住址', 'address'],
    }
    label_lower = label.lower()
    for field_type, keywords in type_keywords.items():
        for keyword in keywords:
            if keyword in label_lower:
                return field_type
    return "text"


def _clean_label(text: str) -> str:
    """清理标签文本：去掉尾部冒号、空白"""
    return re.sub(r'[：:\s]+$', '', text.strip())


# ──────────────────────────── xlsx 解析 ────────────────────────────

def parse_xlsx_structure(file_path: str) -> Dict[str, Any]:
    """解析 xlsx 结构：每个 sheet 的类型、表头、字段（纵向表单时）"""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wb = load_workbook(file_path, read_only=True, data_only=True)
    sheets = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        all_rows = list(ws.iter_rows(values_only=True))

        merged_ranges = []
        try:
            merged_ranges = [str(r) for r in ws.merged_cells.ranges]
        except Exception:
            pass

        max_col = ws.max_column or 0
        non_empty_rows = [
            r for r in all_rows
            if any(v is not None and str(v).strip() for v in r)
        ]

        sheet_info: Dict[str, Any] = {
            "name": sheet_name,
            "max_row": ws.max_row,
            "max_col": max_col,
        }
        if merged_ranges:
            sheet_info["merged_ranges"] = merged_ranges

        # ── 纵向表单探测：列少 + "标签|空值"行占主导 ──
        form_fields = []
        if 0 < max_col <= 4 and non_empty_rows:
            label_value_hits = 0
            for r_idx, row in enumerate(all_rows, start=1):
                cells = list(row)
                first = cells[0] if cells else None
                if first is None or not str(first).strip():
                    continue
                if not isinstance(first, str):
                    continue
                rest = cells[1:]
                rest_non_empty = [v for v in rest if v is not None and str(v).strip()]
                # 标签行：第一列是文本，右侧全空或只有一个值
                if len(rest_non_empty) <= 1:
                    label_value_hits += 1
                    label = _clean_label(str(first))
                    if not label or len(label) > 30:
                        continue
                    # 写入坐标 = 标签右侧第一个单元格
                    value_col = 2
                    value_cell = f"{get_column_letter(value_col)}{r_idx}"
                    current = cells[1] if len(cells) > 1 else None
                    form_fields.append({
                        "label": label,
                        "placeholder_type": "xlsx_cell",
                        "location": {"sheet": sheet_name, "cell": value_cell},
                        "current_value": "" if current is None else str(current),
                        "detected_type": _infer_field_type(label),
                    })

            if label_value_hits >= 3 and label_value_hits >= len(non_empty_rows) * 0.6:
                sheet_info["kind"] = "form"
                sheet_info["fields"] = form_fields
                sheets.append(sheet_info)
                continue

        # ── 数据表格探测：真实表头行 + 数据区 ──
        header_row_idx = detect_xlsx_header_row(all_rows) if all_rows else 1

        title_rows = []
        for r_idx, row in enumerate(all_rows[:header_row_idx - 1], start=1):
            texts = [str(v) for v in row if v is not None and str(v).strip()]
            if texts:
                title_rows.append({"row": r_idx, "text": " ".join(texts)})

        headers = []
        if all_rows and header_row_idx <= len(all_rows):
            header_row = all_rows[header_row_idx - 1]
            headers = [str(cell) if cell is not None and str(cell).strip()
                       else f"Column_{i + 1}"
                       for i, cell in enumerate(header_row)]

        data_rows = all_rows[header_row_idx:] if all_rows else []
        sample_data = []
        for row in data_rows[:3]:
            row_data = {}
            for header, cell in zip(headers, row):
                row_data[header] = str(cell) if cell is not None else ""
            if any(row_data.values()):
                sample_data.append(row_data)

        non_empty_headers = sum(1 for h in headers if not h.startswith("Column_"))
        sheet_info.update({
            "kind": "data_table" if non_empty_headers >= 2 else "unknown",
            "header_row_index": header_row_idx,
            "headers": headers,
            "column_count": len(headers),
            "row_count": len(data_rows),
            "sample_data": sample_data,
        })
        if title_rows:
            sheet_info["title_rows"] = title_rows
        if header_row_idx > 1:
            sheet_info["note"] = (
                f"检测到表头在第{header_row_idx}行（非第1行），"
                f"填表时请使用 header_row={header_row_idx} 参数"
            )
        sheets.append(sheet_info)

    wb.close()
    return {"sheets": sheets}


# ──────────────────────────── docx 解析 ────────────────────────────

def parse_docx_tables(doc) -> List[Dict[str, Any]]:
    """解析 docx 中所有表格的结构（表头、示例、上下文）"""
    table_contexts = _extract_table_contexts(doc)
    tables = []

    for table_idx, table in enumerate(doc.tables):
        headers = []
        if table.rows:
            first_row = table.rows[0]
            headers = [cell.text.strip() if cell.text.strip() else f"Column_{i + 1}"
                       for i, cell in enumerate(first_row.cells)]

        row_count = len(table.rows) - 1 if len(table.rows) > 1 else 0

        sample_data = []
        for row in table.rows[1:4]:
            row_data = {}
            for header, cell in zip(headers, row.cells):
                row_data[header] = cell.text.strip()
            if any(row_data.values()):
                sample_data.append(row_data)

        # 表格用途启发式分类
        non_empty_headers = sum(1 for h in headers if not h.startswith("Column_"))
        data_rows = list(table.rows[1:])
        empty_data_rows = sum(
            1 for row in data_rows
            if all(not cell.text.strip() for cell in row.cells)
        )
        label_value_rows = sum(
            1 for row in data_rows
            if row.cells and row.cells[0].text.strip()
            and all(not cell.text.strip() for cell in row.cells[1:])
        )

        if data_rows and label_value_rows >= len(data_rows) * 0.6:
            kind = "form_table"
        elif non_empty_headers >= 2:
            kind = "data_table"
        else:
            kind = "unknown"

        tables.append({
            "index": table_idx,
            "kind": kind,
            "headers": headers,
            "column_count": len(headers),
            "row_count": row_count,
            "empty_data_rows": empty_data_rows,
            "sample_data": sample_data,
            "context": table_contexts.get(table_idx, {}),
        })

    return tables


def _extract_table_contexts(doc) -> Dict[int, Dict[str, Any]]:
    """提取每个表格的上下文信息（前面的段落文本）"""
    contexts = {}
    table_idx = 0
    prev_paragraphs = []

    for element in doc.element.body:
        if element.tag.endswith('tbl'):
            context_text = []
            for p in prev_paragraphs[-5:]:
                text = p.text.strip()
                if text:
                    context_text.append(text)
            contexts[table_idx] = {"preceding_text": context_text}
            table_idx += 1
            prev_paragraphs = []
        elif element.tag.endswith('p'):
            for p in doc.paragraphs:
                if p._element is element:
                    prev_paragraphs.append(p)
                    break

    return contexts


def scan_paragraph_placeholders(doc) -> List[Dict[str, Any]]:
    """启发式扫描段落占位符（下划线/方括号/行尾冒号），无 LLM

    location.paragraph_index 使用【全部段落】索引（含空段落），
    与 fill_form 的填写实现约定一致（与编辑工具的非空段落索引不同，互不使用）。
    """
    fields = []

    for para_idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue

        if re.search(r'_{3,}', text):
            fields.append({
                "label": _extract_label_from_text(text),
                "placeholder_type": "underline",
                "location": {"paragraph_index": para_idx},
                "current_value": "",
                "detected_type": "text",
            })
        elif re.search(r'【\s*】|\[\s*\]|（\s*）', text):
            fields.append({
                "label": _extract_label_from_text(text),
                "placeholder_type": "bracket",
                "location": {"paragraph_index": para_idx},
                "current_value": "",
                "detected_type": "text",
            })
        elif re.search(r'[：:]\s*$', text):
            label = _clean_label(re.split(r'[：:]', text)[0])
            if label and len(label) <= 20:
                fields.append({
                    "label": label,
                    "placeholder_type": "colon",
                    "location": {"paragraph_index": para_idx},
                    "current_value": "",
                    "detected_type": _infer_field_type(label),
                })

    return fields


def _extract_label_from_text(text: str) -> str:
    """从文本中提取标签"""
    label = _clean_label(re.split(r'[：:]', text)[0])
    if label and len(label) <= 20:
        return label
    label = text.split()[0] if text.split() else text
    return label[:20] if label else "未知字段"


# ──────────────────────────── docx 表单字段 LLM 分块检测 ────────────────────────────

CHUNK_MAX_CHARS = 8000  # 每个分析块的最大字符数


async def detect_form_fields_llm(doc, heuristic_fields: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """LLM 分块分析表单字段（长文档自动分块后合并），失败回退到启发式结果"""
    from app.services.llm_service import llm_service

    lines = _extract_document_lines(doc)
    chunks = _split_into_chunks(lines)

    all_fields: List[Dict[str, Any]] = []
    failed_chunks = 0

    for chunk_lines in chunks:
        chunk_content = "\n".join(chunk_lines)
        fields = await _analyze_chunk_with_llm(chunk_content, llm_service)
        if fields is None:
            failed_chunks += 1
            continue
        all_fields.extend(fields)

    if failed_chunks == len(chunks):
        # 所有块都失败，回退到启发式检测
        return heuristic_fields

    # 按 (label, location) 去重（块边界处可能重复检测）
    seen = set()
    unique_fields = []
    for field in all_fields:
        key = (field["label"], json.dumps(field["location"], sort_keys=True))
        if key in seen:
            continue
        seen.add(key)
        unique_fields.append(field)

    # 启发式检测到但 LLM 漏掉的字段补进来（占位符正则非常可靠）
    for field in heuristic_fields:
        key = (field["label"], json.dumps(field["location"], sort_keys=True))
        if key not in seen:
            seen.add(key)
            unique_fields.append(field)

    return unique_fields


def _split_into_chunks(lines: List[str]) -> List[List[str]]:
    """按字符数将内容行切成多个块（不截断单行）"""
    chunks: List[List[str]] = []
    current: List[str] = []
    current_len = 0

    for line in lines:
        line_len = len(line) + 1
        if current and current_len + line_len > CHUNK_MAX_CHARS:
            chunks.append(current)
            current = []
            current_len = 0
        current.append(line)
        current_len += line_len

    if current:
        chunks.append(current)

    return chunks or [[]]


def _extract_document_lines(doc) -> List[str]:
    """提取文档内容行供LLM分析（段落/表格均带全局索引标签）

    段落索引为【全部段落】索引（含空段落不输出但占编号），
    与 fill_form 的填写实现约定一致。
    """
    content_parts = []

    for para_idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            content_parts.append(f"[段落{para_idx}] {text}")

    for table_idx, table in enumerate(doc.tables):
        content_parts.append(f"\n[表格{table_idx}]")
        for row_idx, row in enumerate(table.rows):
            cells = []
            for col_idx, cell in enumerate(row.cells):
                cell_text = cell.text.strip()
                cells.append(f"[{col_idx}]{cell_text if cell_text else '(空)'}")
            content_parts.append(f"  行{row_idx}: {' | '.join(cells)}")

    return content_parts


async def _analyze_chunk_with_llm(doc_content: str, llm_service) -> Optional[List[Dict[str, Any]]]:
    """分析单个内容块，返回字段列表；失败返回 None"""
    prompt = f"""请分析以下Word文档内容（可能是完整文档或文档的一部分），识别所有需要填写的字段。

文档内容：
{doc_content}

请以JSON格式返回分析结果，包含：
- fields: 需要填写的字段列表，每个字段包含：
  - type: "paragraph" 或 "table"
  - label: 字段标签（如"姓名"、"实习名称"、"题目"等）
  - location: 位置信息
    - 对于段落：{{"paragraph_index": 段落索引}}
    - 对于表格：{{"table_index": 表格索引, "row_index": 行索引, "cell_index": 列索引}}
  - current_value: 当前内容（如果是示例/说明文本，标记为示例）
  - is_example: 是否是示例/说明文本（应该被覆盖）

判断规则：
1. 段落中的占位符：
   - 下划线（______）后面需要填写
   - 方括号（【】、[]）内需要填写
   - 冒号（：或:）后面没有内容或只有空格，需要填写
   - "请填写"、"请输入"等提示后面需要填写

2. 表格中的字段：
   - 空单元格需要填写
   - 包含示例文本（如"1、xxx 2、xxx"、"如xxx"、"例如xxx"）的单元格需要填写
   - 包含说明性文本（如"数据、算法、算力、人员等可行性"）的单元格需要填写
   - 表头行不需要填写

3. 特殊情况：
   - 如果文档是手册/指南，可能包含多个需要填写的部分
   - 如果文档是审批表，可能包含多个表格
   - 如果文档是申请表，可能包含段落和表格的混合

注意：段落索引和表格索引必须使用内容中的[段落N]/[表格N]标签原样返回，不要重新编号。
请只返回JSON，不要有其他内容。"""

    try:
        response = await llm_service.chat_completion(
            messages=[{"role": "user", "content": prompt}],
            temperature=0.1,
            enable_thinking=False
        )

        analysis = None
        if isinstance(response, str):
            json_match = re.search(r'\{[\s\S]*\}', response)
            if json_match:
                analysis = json.loads(json_match.group())
        elif isinstance(response, dict):
            analysis = response

        if not analysis or "fields" not in analysis:
            return []

        fields = []
        for field_info in analysis.get("fields") or []:
            field_type = field_info.get("type", "paragraph")
            label = field_info.get("label", "")
            location = field_info.get("location", {})
            current_value = field_info.get("current_value", "")
            is_example = field_info.get("is_example", False)

            if not label or not location:
                continue

            fields.append({
                "label": label,
                "placeholder_type": "table_cell" if field_type == "table" else "llm_detected",
                "location": location,
                "current_value": current_value,
                "is_example": bool(is_example),
                "detected_type": _infer_field_type(label),
            })

        return fields

    except Exception as e:
        print(f"LLM分析文档块失败: {e}")
        return None


# ──────────────────────────── 顶层结构分析（供工具与 fill_form 复用） ────────────────────────────

async def analyze_template(doc: Document) -> Dict[str, Any]:
    """分析模板结构（带缓存）。doc 为 Document ORM 对象。

    返回结构：
    - structure_type: data_table / form / mixed / unknown
    - xlsx: sheets[...]（data_table sheet 含 header_row_index/headers；form sheet 含 fields）
    - docx: tables[...] + fields[...]（form/mixed 时）
    - 所有字段统一编 field_id: F1..Fn
    """
    cache_key = structure_cache.make_key("template_structure", str(doc.id), doc.sha256)
    cached = structure_cache.get(cache_key)
    if cached is not None:
        return cached

    if doc.file_type == "xlsx":
        result = _analyze_xlsx(doc.file_path)
    elif doc.file_type == "docx":
        result = await _analyze_docx(doc.file_path)
    else:
        result = {
            "structure_type": "unknown",
            "error": f"不支持的模板格式: {doc.file_type}（仅支持 docx/xlsx）",
        }

    # 统一编 field_id
    fields = result.get("fields") or []
    for i, field in enumerate(fields, start=1):
        field["field_id"] = f"F{i}"
    result["total_fields"] = len(fields)

    structure_cache.set(cache_key, result)
    return result


def _analyze_xlsx(file_path: str) -> Dict[str, Any]:
    parsed = parse_xlsx_structure(file_path)
    sheets = parsed["sheets"]

    kinds = {s.get("kind") for s in sheets}
    has_data = "data_table" in kinds
    has_form = "form" in kinds

    if has_data and has_form:
        structure_type = "mixed"
    elif has_data:
        structure_type = "data_table"
    elif has_form:
        structure_type = "form"
    else:
        structure_type = "unknown"

    # 汇总所有 form sheet 的字段
    fields: List[Dict[str, Any]] = []
    for s in sheets:
        fields.extend(s.get("fields") or [])

    return {
        "structure_type": structure_type,
        "sheets": sheets,
        "fields": fields,
    }


async def _analyze_docx(file_path: str) -> Dict[str, Any]:
    from docx import Document as DocxDocument

    doc = DocxDocument(file_path)

    tables = parse_docx_tables(doc)
    heuristic_fields = scan_paragraph_placeholders(doc)

    has_data_table = any(t["kind"] == "data_table" for t in tables)
    has_form_table = any(t["kind"] == "form_table" for t in tables)
    has_para_fields = bool(heuristic_fields)

    if has_data_table and not has_form_table and not has_para_fields:
        structure_type = "data_table"
    elif has_data_table and (has_form_table or has_para_fields):
        structure_type = "mixed"
    elif has_form_table or has_para_fields:
        structure_type = "form"
    else:
        structure_type = "unknown"

    # 表单字段：有表单迹象时用 LLM 分块检测（含表格内字段），否则不花 LLM 调用
    fields: List[Dict[str, Any]] = []
    form_type = None
    if structure_type in ("form", "mixed", "unknown"):
        fields = await detect_form_fields_llm(doc, heuristic_fields)
        has_table_fields = any(f['location'].get('table_index') is not None for f in fields)
        has_p_fields = any(f['location'].get('paragraph_index') is not None for f in fields)
        if has_p_fields and has_table_fields:
            form_type = "mixed"
        elif has_table_fields:
            form_type = "table_form"
        elif has_p_fields:
            form_type = "unstructured"
        # 修正 unknown：LLM 找到字段说明确实是表单
        if structure_type == "unknown" and fields:
            structure_type = "form"

    result: Dict[str, Any] = {
        "structure_type": structure_type,
        "tables": tables,
        "fields": fields,
    }
    if form_type:
        result["form_type"] = form_type
    return result


# ──────────────────────────── 工具类 ────────────────────────────

class GetTemplateStructureTool(BaseTool):
    """模板结构分析工具

    填表/填表单前的第一步：一次调用返回模板类型 + 表格结构 + 表单字段。
    """

    @property
    def name(self) -> str:
        return "get_template_structure"

    @property
    def description(self) -> str:
        return """分析模板文档结构，填表/填表单一站式探测（填写前必须先调用）。

返回内容：
- structure_type: data_table(数据表格) / form(表单) / mixed(混合) / unknown
- xlsx 数据表格: 每个 sheet 的表头（自动探测真实表头行，支持标题行/说明行）、
  header_row_index、示例数据、合并单元格；多工作表全部列出
- xlsx 纵向表单（A列标签B列填值）: 每个字段的标签和写入坐标（fields[].location.cell）
- docx 数据表格: 每个表格的表头、行列数、示例数据、context.preceding_text（表格用途）
- docx 表单: 所有可填写字段 fields[]，每个字段带稳定 field_id（F1..Fn）、
  label、placeholder_type、location、detected_type

后续动作指引：
- structure_type=data_table → 用 fill_table 填写（xlsx 多工作表传 sheet_name；
  header_row_index>1 时传 header_row；docx 多表格传 target_table_index）
- structure_type=form → 用 fill_form 填写，fields 参数引用 field_id
- mixed → 两种工具按需各填一部分

同一份模板重复调用会命中缓存（内容变化后自动重新分析），
fill_form 也会复用本工具的字段检测结果，无需担心字段对不上。"""

    @property
    def parameters(self) -> Dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "template_id": {
                    "type": "string",
                    "description": "模板文档ID"
                },
                "refresh": {
                    "type": "boolean",
                    "default": False,
                    "description": "强制重新分析（忽略缓存），默认false"
                }
            },
            "required": ["template_id"]
        }

    async def execute(self, params: Dict[str, Any], context: ToolContext) -> ToolResult:
        try:
            template_id = params.get("template_id", "")
            uuid_error = BaseTool.validate_uuid(template_id, "template_id")
            if uuid_error:
                return ToolResult(success=False, error=uuid_error)

            async with async_session() as db:
                result = await db.execute(
                    select(Document).where(Document.id == template_id)
                )
                doc = result.scalar_one_or_none()

            if not doc:
                return ToolResult(success=False, error=f"模板文档不存在: {template_id}")

            if params.get("refresh"):
                cache_key = structure_cache.make_key("template_structure", str(doc.id), doc.sha256)
                structure_cache.invalidate(cache_key)

            analysis = await analyze_template(doc)

            return ToolResult(
                success=True,
                data={
                    "template_id": template_id,
                    "filename": doc.original_filename,
                    "file_type": doc.file_type,
                    **analysis,
                }
            )

        except Exception as e:
            return ToolResult(success=False, error=f"获取模板结构失败: {str(e)}")
