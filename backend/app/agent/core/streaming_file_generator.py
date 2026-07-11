"""
流式文件生成器 - 支持边生成边返回进度

负责：
- 分块写入数据
- 流式进度报告
"""

from typing import Dict, Any, List
import logging

logger = logging.getLogger(__name__)


class StreamingFileGenerator:
    """流式文件生成器

    支持边生成边返回进度。
    """

    CHUNK_SIZE = 50  # 每块行数

    async def generate_and_stream(
        self,
        template_id: str,
        data: List[Dict[str, Any]],
        mapping: Dict[str, str],
        stream_bus: Any,
        task_id: str
    ) -> Dict[str, Any]:
        """生成文件并流式报告进度

        Args:
            template_id: 模板ID
            data: 数据
            mapping: 列映射
            stream_bus: 流总线
            task_id: 任务ID

        Returns:
            文件信息
        """
        logger.info(f"[StreamingFileGenerator] 开始生成文件 | 数据行数: {len(data)}")

        # 分块写入
        total_rows = len(data)
        written_rows = 0

        # 获取模板信息
        template_info = await self._get_template_info(template_id)

        # 创建输出文件
        output_path = await self._create_output_file(template_info)

        # 分块写入数据
        for i in range(0, total_rows, self.CHUNK_SIZE):
            chunk = data[i:i + self.CHUNK_SIZE]

            # 写入数据块
            await self._write_chunk(output_path, chunk, mapping)

            # 更新进度
            written_rows += len(chunk)
            progress = written_rows / total_rows

            # 发送进度事件
            if stream_bus:
                await stream_bus.emit_progress(
                    task_id=task_id,
                    progress=progress,
                    message=f"已填写 {written_rows}/{total_rows} 行"
                )

            logger.info(f"[StreamingFileGenerator] 进度: {progress:.1%} ({written_rows}/{total_rows})")

        # 上传文件
        file_info = await self._upload_file(output_path, template_info)

        logger.info(f"[StreamingFileGenerator] 文件生成完成 | {file_info}")

        return file_info

    async def _get_template_info(self, template_id: str) -> Dict[str, Any]:
        """获取模板信息"""
        from app.db.postgres import async_session
        from app.models.document import Document
        from sqlalchemy import select

        async with async_session() as db:
            result = await db.execute(
                select(Document).where(Document.id == template_id)
            )
            doc = result.scalar_one_or_none()

            if not doc:
                raise ValueError(f"模板不存在: {template_id}")

            return {
                "id": str(doc.id),
                "filename": doc.original_filename,
                "file_path": doc.file_path,
                "file_type": doc.file_type
            }

    async def _create_output_file(self, template_info: Dict[str, Any]) -> str:
        """创建输出文件"""
        import os
        import shutil
        from uuid import uuid4

        # 创建输出目录
        output_dir = os.path.join(os.path.dirname(template_info["file_path"]), "outputs")
        os.makedirs(output_dir, exist_ok=True)

        # 复制模板
        output_filename = f"fast_{uuid4().hex[:8]}_{template_info['filename']}"
        output_path = os.path.join(output_dir, output_filename)

        shutil.copy2(template_info["file_path"], output_path)

        return output_path

    async def _write_chunk(
        self,
        output_path: str,
        chunk: List[Dict[str, Any]],
        mapping: Dict[str, str]
    ):
        """写入数据块"""
        try:
            from openpyxl import load_workbook

            wb = load_workbook(output_path)
            ws = wb.active

            # 获取当前行数
            current_row = ws.max_row + 1

            # 写入数据
            for row_data in chunk:
                for col_name, value in row_data.items():
                    # 使用映射找到对应的列
                    mapped_col = mapping.get(col_name, col_name)

                    # 查找列索引
                    col_index = self._find_column_index(ws, mapped_col)
                    if col_index:
                        ws.cell(row=current_row, column=col_index, value=value)

                current_row += 1

            wb.save(output_path)
            wb.close()

        except Exception as e:
            logger.error(f"写入数据块失败: {e}")
            raise

    def _find_column_index(self, ws, col_name: str) -> int:
        """查找列索引"""
        for col in range(1, ws.max_column + 1):
            if ws.cell(row=1, column=col).value == col_name:
                return col
        return 0

    async def _upload_file(self, output_path: str, template_info: Dict[str, Any]) -> Dict[str, Any]:
        """上传文件"""
        import os

        # 注册文件到数据库
        from app.db.postgres import async_session
        from app.models.document import Document

        output_filename = os.path.basename(output_path)

        async with async_session() as db:
            output_doc = Document(
                filename=output_filename,
                original_filename=output_filename,
                file_path=output_path,
                file_type=template_info["file_type"],
                doc_category="output",
                status="completed",
                file_size=os.path.getsize(output_path)
            )
            db.add(output_doc)
            await db.commit()
            await db.refresh(output_doc)

            return {
                "file_id": str(output_doc.id),
                "filename": output_filename,
                "download_url": f"/api/v1/documents/{output_doc.id}/download"
            }


# 全局流式文件生成器实例
streaming_file_generator = StreamingFileGenerator()
