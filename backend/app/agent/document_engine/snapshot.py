"""
文档快照：一次打开 docx/xlsx 到内存，供一组原子操作复用

命名隔离约定：
- snapshot 是会话内的内存概念（本模块）
- version 是持久化概念（app.services.document_versioning，DB 行 + 磁盘文件）
两者永不同名出现。

当前策略：每次 commit 重新打开文件（简单、正确）。
若后续出现性能问题，再引入 run 级快照缓存。
"""

import logging
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


class DocSnapshot:
    """文档快照：持有 python-docx Document 或 openpyxl Workbook 之一"""

    def __init__(self, path: str, file_type: str, docx=None, wb=None):
        self.path = str(path)
        self.file_type = file_type
        self.docx = docx
        self.wb = wb

    @classmethod
    def open(cls, path: str, file_type: Optional[str] = None) -> "DocSnapshot":
        """打开文件为快照。file_type 缺省时按扩展名推断。"""
        p = Path(path)
        if not p.exists():
            raise FileNotFoundError(f"文件不存在: {path}")
        ftype = (file_type or p.suffix.lstrip(".")).lower()

        if ftype == "docx":
            from docx import Document as DocxDocument
            return cls(str(p), "docx", docx=DocxDocument(str(p)))
        if ftype == "xlsx":
            from openpyxl import load_workbook
            return cls(str(p), "xlsx", wb=load_workbook(str(p)))
        raise ValueError(f"快照不支持的文件类型: {ftype}（支持 docx/xlsx；md/txt 按文本行处理，不经快照）")

    def save(self, path: Optional[str] = None) -> None:
        """写回磁盘（默认写回原路径）"""
        target = path or self.path
        if self.docx is not None:
            self.docx.save(target)
        elif self.wb is not None:
            self.wb.save(target)

    def close(self) -> None:
        if self.wb is not None:
            try:
                self.wb.close()
            except Exception:
                pass
            self.wb = None
        self.docx = None

    def __enter__(self) -> "DocSnapshot":
        return self

    def __exit__(self, *exc) -> None:
        self.close()
