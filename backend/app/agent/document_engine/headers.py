"""
表头解析引擎（从 fill_table_tool 抽取）

- detect_xlsx_header_row: 探测真实表头行（处理标题行/说明行）
- resolve_xlsx_headers: 解析指定 sheet 的表头
- get_template_headers: xlsx/docx 统一入口
"""

from typing import List, Optional, Tuple


def detect_xlsx_header_row(rows: List[Tuple], max_scan: int = 10) -> int:
    """探测 Excel 的真实表头行（处理标题行、说明行等情况）。

    打分规则：
    - 非空率高的行更可能是表头
    - 文本比例高的行更可能是表头（数据行通常含数值）
    - 下一行非空数与当前行接近的，更可能是"表头→数据"的衔接

    Returns:
        表头行号（1-based）
    """
    best_row, best_score = 1, -1.0

    for i, row in enumerate(rows[:max_scan], start=1):
        cells = list(row)
        non_empty = [v for v in cells if v is not None and str(v).strip() != ""]
        if not non_empty:
            continue

        fill_rate = len(non_empty) / max(len(cells), 1)
        text_rate = sum(1 for v in non_empty if isinstance(v, str)) / len(non_empty)

        next_bonus = 0.0
        if i < len(rows):
            next_non_empty = sum(
                1 for v in rows[i] if v is not None and str(v).strip() != ""
            )
            if next_non_empty >= len(non_empty) * 0.5:
                next_bonus = 1.0

        score = fill_rate * 2 + text_rate + next_bonus
        if score > best_score:
            best_score, best_row = score, i

    return best_row


def resolve_xlsx_headers(
    file_path: str, sheet_name: Optional[str] = None, header_row: Optional[int] = None
) -> Tuple[List[str], int]:
    """解析 Excel 的表头，返回 (headers, resolved_header_row)。

    header_row 未指定时自动探测真实表头行（处理标题行/说明行）。
    """
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"工作表不存在: {sheet_name}，可用工作表: {list(wb.sheetnames)}")
            ws = wb[sheet_name]
        else:
            ws = wb[wb.sheetnames[0]]

        if header_row:
            resolved_row = header_row
            header_cells = next(
                ws.iter_rows(min_row=resolved_row, max_row=resolved_row, values_only=True),
                None,
            )
        else:
            all_rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                all_rows.append(row)
                if i >= 9:  # 探测只需前10行
                    break
            resolved_row = detect_xlsx_header_row(all_rows) if all_rows else 1
            header_cells = all_rows[resolved_row - 1] if all_rows else None

        headers = []
        if header_cells:
            headers = [
                str(cell) if cell is not None and str(cell).strip() else f"Column_{i + 1}"
                for i, cell in enumerate(header_cells)
            ]
        return headers, resolved_row
    finally:
        wb.close()


def get_template_headers(
    file_path: str,
    file_type: str,
    target_table_index: Optional[int] = None,
    sheet_name: Optional[str] = None,
    header_row: Optional[int] = None,
) -> List[str]:
    """获取模板文件的表头（xlsx 按 sheet，docx 按表格第一行）"""
    if file_type == "xlsx":
        headers, _ = resolve_xlsx_headers(file_path, sheet_name, header_row)
        return headers
    if file_type == "docx":
        from docx import Document as DocxDocument

        doc = DocxDocument(file_path)
        if not doc.tables:
            return []
        table_idx = target_table_index if target_table_index is not None else 0
        table = doc.tables[min(table_idx, len(doc.tables) - 1)]
        if table.rows:
            return [
                cell.text.strip() if cell.text.strip() else f"Column_{i + 1}"
                for i, cell in enumerate(table.rows[0].cells)
            ]
    return []


def match_value_for_header(row_data: dict, header: str) -> str:
    """三级列名匹配兜底：精确 → 大小写不敏感 → 包含

    主映射应由 LLM 显式给出（column_map）；本函数仅作兜底。
    跳过 `_source`/`chunk` 等内部保留键，防止包含匹配误把标签写进表格。
    """
    from app.agent.document_engine.provenance import RESERVED_KEYS

    if header in row_data and row_data[header] is not None:
        return str(row_data[header])
    header_lower = header.strip().lower()
    for key, value in row_data.items():
        if key in RESERVED_KEYS:
            continue
        if value is not None and key.strip().lower() == header_lower:
            return str(value)
    for key, value in row_data.items():
        if key in RESERVED_KEYS:
            continue
        if value is not None:
            key_clean = key.strip().lower()
            if header_lower in key_clean or key_clean in header_lower:
                return str(value)
    return ""
