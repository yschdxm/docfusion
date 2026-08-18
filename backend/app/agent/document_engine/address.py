"""
地址模型与地址解析器

设计原则：快照 + 内容锚点优先 + 索引兜底。
- docx 没有可靠的段落持久 ID，不设计"稳定 ID"体系
- 锚点（段落文本前20字）优先用于校验/定位，索引作为获取手段与兜底
- 解析结果三态：ok / ambiguous（锚点重复）/ not_found

地址是会话内快照概念，与 document_versioning 的持久 version 无关。
"""

from typing import Any, List, Literal, Optional, Union

from pydantic import BaseModel, Field


# ──────────────────────────── 地址模型 ────────────────────────────

class XlsxCellAddr(BaseModel):
    """Excel 单元格地址：sheet + A1 坐标"""
    kind: Literal["xlsx_cell"] = "xlsx_cell"
    cell: str = Field(..., description="A1 风格坐标，如 B5")
    sheet: Optional[str] = Field(None, description="工作表名，缺省为活动工作表")


class XlsxSheetAddr(BaseModel):
    """Excel 工作表地址（统计表填写的目标）"""
    kind: Literal["xlsx_sheet"] = "xlsx_sheet"
    sheet: Optional[str] = Field(None, description="工作表名，缺省为活动工作表")
    header_row: Optional[int] = Field(None, description="表头行号（1-based），缺省自动探测")


class DocxTableCellAddr(BaseModel):
    """Word 表格单元格地址（0-based）"""
    kind: Literal["docx_table_cell"] = "docx_table_cell"
    table_index: int = Field(..., description="表格索引（从0开始）")
    row: int = Field(..., description="行索引（从0开始，第0行通常是表头）")
    col: int = Field(..., description="列索引（从0开始）")


class DocxTableAddr(BaseModel):
    """Word 表格地址（统计表填写的目标）"""
    kind: Literal["docx_table"] = "docx_table"
    table_index: int = Field(..., description="表格索引（从0开始）")


class DocxParagraphAddr(BaseModel):
    """Word 段落地址：索引 + 内容锚点

    index_scope:
    - all: 全部段落索引（含空段落），与 fill_form 的 location 约定一致
    - editable: 非空段落索引，与 get_document_outline 的 [Pn] 一致
    """
    kind: Literal["docx_paragraph"] = "docx_paragraph"
    paragraph_index: Optional[int] = Field(None, description="段落索引（从0开始）")
    anchor: Optional[str] = Field(None, description="段落文本前20字，用于校验/抗漂移定位")
    index_scope: Literal["all", "editable"] = "all"


class ContentControlAddr(BaseModel):
    """Word 内容控件地址：tag 优先，sdt_index 兜底"""
    kind: Literal["content_control"] = "content_control"
    tag: Optional[str] = Field(None, description="控件 w:tag（稳定标识）")
    sdt_index: Optional[int] = Field(None, description="控件枚举索引（无 tag 时兜底）")


Address = Union[XlsxCellAddr, XlsxSheetAddr, DocxTableCellAddr, DocxTableAddr, DocxParagraphAddr, ContentControlAddr]


# ──────────────────────────── 解析结果 ────────────────────────────

class Resolved(BaseModel):
    """地址解析结果。target 为 python-docx/openpyxl 原生对象。"""
    ok: bool
    status: Literal["ok", "ambiguous", "not_found"] = "ok"
    target: Any = None
    detail: str = ""

    class Config:
        arbitrary_types_allowed = True


def describe_address(addr: Address) -> str:
    """地址的人类可读描述（用于报告与日志）"""
    if isinstance(addr, XlsxCellAddr):
        return f"{addr.sheet or '活动表'}!{addr.cell}"
    if isinstance(addr, XlsxSheetAddr):
        return f"工作表[{addr.sheet or '活动表'}]"
    if isinstance(addr, DocxTableCellAddr):
        return f"表格{addr.table_index} R{addr.row}C{addr.col}"
    if isinstance(addr, DocxTableAddr):
        return f"表格{addr.table_index}"
    if isinstance(addr, DocxParagraphAddr):
        idx = f"P{addr.paragraph_index}" if addr.paragraph_index is not None else "P?"
        anchor = f"「{addr.anchor[:15]}…」" if addr.anchor else ""
        return f"段落{idx}{anchor}"
    if isinstance(addr, ContentControlAddr):
        return f"内容控件[tag={addr.tag}]" if addr.tag else f"内容控件[{addr.sdt_index}]"
    return str(addr)


# ──────────────────────────── 解析器 ────────────────────────────

class AddressResolver:
    """把 Address 解析为文档对象。解析不写盘、不修改文档。"""

    @staticmethod
    def editable_paragraphs(doc) -> list:
        """非空段落列表（原 _get_editable_paragraphs 逻辑的单点维护处）"""
        return [p for p in doc.paragraphs if p.text.strip()]

    @classmethod
    def resolve(cls, addr: Address, snapshot) -> Resolved:
        if isinstance(addr, XlsxCellAddr):
            return cls._resolve_xlsx_cell(addr, snapshot)
        if isinstance(addr, XlsxSheetAddr):
            return cls._resolve_xlsx_sheet(addr, snapshot)
        if isinstance(addr, DocxTableCellAddr):
            return cls._resolve_docx_table_cell(addr, snapshot)
        if isinstance(addr, DocxTableAddr):
            return cls._resolve_docx_table(addr, snapshot)
        if isinstance(addr, DocxParagraphAddr):
            return cls._resolve_docx_paragraph(addr, snapshot)
        if isinstance(addr, ContentControlAddr):
            return cls._resolve_content_control(addr, snapshot)
        return Resolved(ok=False, status="not_found", detail=f"未知地址类型: {type(addr).__name__}")

    # ── xlsx ──

    @staticmethod
    def _resolve_xlsx_cell(addr: XlsxCellAddr, snapshot) -> Resolved:
        from openpyxl.utils import coordinate_to_tuple
        from openpyxl.utils.exceptions import InvalidFileException

        wb = snapshot.wb
        if addr.sheet:
            if addr.sheet not in wb.sheetnames:
                return Resolved(ok=False, status="not_found",
                                detail=f"工作表不存在: {addr.sheet}，可用: {wb.sheetnames}")
            ws = wb[addr.sheet]
        else:
            ws = wb.active
        try:
            row, col = coordinate_to_tuple(addr.cell.upper())
        except (ValueError, InvalidFileException):
            return Resolved(ok=False, status="not_found", detail=f"非法单元格坐标: {addr.cell}")
        return Resolved(ok=True, target=ws.cell(row=row, column=col))

    @staticmethod
    def _resolve_xlsx_sheet(addr: XlsxSheetAddr, snapshot) -> Resolved:
        wb = snapshot.wb
        if addr.sheet:
            if addr.sheet not in wb.sheetnames:
                return Resolved(ok=False, status="not_found",
                                detail=f"工作表不存在: {addr.sheet}，可用: {wb.sheetnames}")
            return Resolved(ok=True, target=wb[addr.sheet])
        return Resolved(ok=True, target=wb.active)

    # ── docx 表格 ──

    @staticmethod
    def _resolve_docx_table_cell(addr: DocxTableCellAddr, snapshot) -> Resolved:
        doc = snapshot.docx
        if addr.table_index >= len(doc.tables):
            return Resolved(ok=False, status="not_found",
                            detail=f"表格索引 {addr.table_index} 超出范围（共 {len(doc.tables)} 个表格）")
        table = doc.tables[addr.table_index]
        if addr.row >= len(table.rows):
            return Resolved(ok=False, status="not_found",
                            detail=f"行索引 {addr.row} 超出范围（表格共 {len(table.rows)} 行）")
        row = table.rows[addr.row]
        if addr.col >= len(row.cells):
            return Resolved(ok=False, status="not_found",
                            detail=f"列索引 {addr.col} 超出范围（该行共 {len(row.cells)} 列）")
        return Resolved(ok=True, target=row.cells[addr.col])

    @staticmethod
    def _resolve_docx_table(addr: DocxTableAddr, snapshot) -> Resolved:
        doc = snapshot.docx
        if addr.table_index >= len(doc.tables):
            return Resolved(ok=False, status="not_found",
                            detail=f"表格索引 {addr.table_index} 超出范围（共 {len(doc.tables)} 个表格）")
        return Resolved(ok=True, target=doc.tables[addr.table_index])

    # ── docx 段落（锚点优先校验，索引获取，锚点重定位兜底） ──

    @classmethod
    def _resolve_docx_paragraph(cls, addr: DocxParagraphAddr, snapshot) -> Resolved:
        doc = snapshot.docx
        paragraphs = doc.paragraphs if addr.index_scope == "all" else cls.editable_paragraphs(doc)

        anchor = (addr.anchor or "").strip()

        # 1. 索引直达 + 锚点校验
        if addr.paragraph_index is not None:
            if 0 <= addr.paragraph_index < len(paragraphs):
                para = paragraphs[addr.paragraph_index]
                if not anchor or cls._anchor_matches(anchor, para.text):
                    return Resolved(ok=True, target=para)
                # 锚点不匹配 → 索引可能已漂移，尝试锚点重定位
            elif not anchor:
                return Resolved(ok=False, status="not_found",
                                detail=f"段落索引 {addr.paragraph_index} 超出范围（共 {len(paragraphs)} 个）")

        # 2. 锚点定位
        if anchor:
            matches = [p for p in paragraphs if cls._anchor_matches(anchor, p.text)]
            if len(matches) == 1:
                return Resolved(ok=True, target=matches[0],
                                detail="索引已漂移，通过锚点重定位" if addr.paragraph_index is not None else "")
            if len(matches) > 1:
                return Resolved(ok=False, status="ambiguous",
                                detail=f"锚点「{anchor[:20]}」匹配到 {len(matches)} 个段落，请提供更长的锚点或同时给出索引")
            return Resolved(ok=False, status="not_found",
                            detail=f"未找到锚点「{anchor[:20]}」对应的段落")

        return Resolved(ok=False, status="not_found", detail="段落地址缺少 paragraph_index 和 anchor")

    @staticmethod
    def _anchor_matches(anchor: str, text: str) -> bool:
        """锚点匹配：锚点是文本前缀（文本可能被截断传入，做双向包容）"""
        stripped = text.strip()
        if not stripped:
            return False
        probe = anchor[:20]
        # 正常：文本以锚点开头；包容：短文本被编辑截断后整体仍是锚点前缀
        return stripped.startswith(probe) or (len(stripped) >= 8 and probe.startswith(stripped))

    # ── docx 内容控件 ──

    @staticmethod
    def _resolve_content_control(addr: ContentControlAddr, snapshot) -> Resolved:
        from docx.oxml.ns import qn

        doc = snapshot.docx
        sdt_elements = doc.element.body.findall('.//' + qn('w:sdt'))

        if addr.tag:
            matches = []
            for sdt in sdt_elements:
                tag_el = sdt.find('.//' + qn('w:sdtPr') + '/' + qn('w:tag'))
                if tag_el is None:
                    sdt_pr = sdt.find(qn('w:sdtPr'))
                    tag_el = sdt_pr.find(qn('w:tag')) if sdt_pr is not None else None
                if tag_el is not None and tag_el.get(qn('w:val')) == addr.tag:
                    matches.append(sdt)
            if len(matches) == 1:
                return Resolved(ok=True, target=matches[0])
            if len(matches) > 1:
                return Resolved(ok=False, status="ambiguous",
                                detail=f"tag={addr.tag} 匹配到 {len(matches)} 个内容控件")
            # tag 未命中时若有 sdt_index 则兜底
            if addr.sdt_index is None:
                return Resolved(ok=False, status="not_found",
                                detail=f"未找到 tag={addr.tag} 的内容控件")

        if addr.sdt_index is not None:
            if 0 <= addr.sdt_index < len(sdt_elements):
                return Resolved(ok=True, target=sdt_elements[addr.sdt_index])
            return Resolved(ok=False, status="not_found",
                            detail=f"内容控件索引 {addr.sdt_index} 超出范围（共 {len(sdt_elements)} 个）")

        return Resolved(ok=False, status="not_found", detail="内容控件地址缺少 tag 和 sdt_index")


def list_content_controls(doc) -> List[dict]:
    """枚举文档中的内容控件（供结构分析使用）

    Returns:
        [{sdt_index, tag, alias, text}] —— text 为控件当前纯文本（截断）
    """
    from docx.oxml.ns import qn

    controls = []
    for i, sdt in enumerate(doc.element.body.findall('.//' + qn('w:sdt'))):
        sdt_pr = sdt.find(qn('w:sdtPr'))
        tag = alias = None
        if sdt_pr is not None:
            tag_el = sdt_pr.find(qn('w:tag'))
            alias_el = sdt_pr.find(qn('w:alias'))
            tag = tag_el.get(qn('w:val')) if tag_el is not None else None
            alias = alias_el.get(qn('w:val')) if alias_el is not None else None
        text = "".join(t.text or "" for t in sdt.findall('.//' + qn('w:t')))
        controls.append({
            "sdt_index": i,
            "tag": tag,
            "alias": alias,
            "text": text.strip()[:50],
        })
    return controls
