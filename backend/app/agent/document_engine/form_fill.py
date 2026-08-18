"""
表单填写引擎：字段匹配（mapping 草稿）+ 定位（locator）+ 纯文本变换 + 统一写入

7 种占位符策略收敛为：
  locator（location → Address）→ AddressResolver 解析 → transform（纯函数算新文本）→ ops 写入
prepare_* 返回 PreparedFill，dry_run 时只读 before/after，commit 时调 apply() 落写。
"""

import logging
import re
from difflib import SequenceMatcher
from typing import Any, Callable, Dict, List, Optional

from pydantic import BaseModel

from app.agent.document_engine import ops
from app.agent.document_engine.address import (
    AddressResolver,
    ContentControlAddr,
    DocxParagraphAddr,
    DocxTableCellAddr,
    XlsxCellAddr,
    describe_address,
)
from app.agent.document_engine.snapshot import DocSnapshot

logger = logging.getLogger(__name__)


class PreparedFill(BaseModel):
    """一个字段的填写准备结果：dry_run 读报告，commit 调 apply()"""
    subject: str
    ok: bool
    status: str = "ok"  # ok / not_found / ambiguous
    detail: str = ""
    before: str = ""
    after: str = ""

    class Config:
        arbitrary_types_allowed = True

    # apply 不属于 schema，运行期挂载
    _apply: Optional[Callable[[], "ops.OpResult"]] = None

    def set_apply(self, fn: Callable[[], "ops.OpResult"]) -> None:
        object.__setattr__(self, "_apply", fn)

    def apply(self) -> "ops.OpResult":
        if self._apply is None:
            return ops.OpResult.fail(self.subject, "该字段不可写入（校验未通过）")
        return self._apply()


# ──────────────────────────── 字段匹配（mapping 草稿，纯确定性） ────────────────────────────

def calculate_match_score(data_key: str, field_label: str) -> float:
    """数据key与字段标签的匹配分数"""
    if data_key == field_label:
        return 1.0
    if data_key.lower() == field_label.lower():
        return 0.95
    if data_key in field_label or field_label in data_key:
        return 0.8

    synonyms = {
        '姓名': ['名字', 'name', '用户名', '联系人'],
        '电话': ['手机', '联系方式', 'phone', 'tel', '联系电话'],
        '邮箱': ['邮件', 'email', '电子邮箱'],
        '地址': ['住址', 'address', '通讯地址'],
        '日期': ['时间', 'date', '签约日期'],
        '甲方': ['委托方', '客户方', '买方'],
        '乙方': ['受托方', '服务方', '卖方'],
    }
    for canonical, syn_list in synonyms.items():
        if data_key in [canonical] + syn_list and field_label in [canonical] + syn_list:
            return 0.85

    return SequenceMatcher(None, data_key.lower(), field_label.lower()).ratio()


def resolve_field_items(
    field_items: List[Dict[str, Any]],
    form_fields: List[Dict[str, Any]],
) -> tuple:
    """解析 fields 数组：按 field_id 精确匹配，无 field_id 时按 label 匹配

    Returns:
        (matched_fields, unmatched_keys)
    """
    by_id = {f.get("field_id"): f for f in form_fields if f.get("field_id")}
    matched = []
    unmatched = []
    used_field_ids = set()

    for i, item in enumerate(field_items):
        value = item.get("value")
        field = None
        data_key = None

        field_id = item.get("field_id")
        if field_id:
            data_key = field_id
            field = by_id.get(field_id)
            if field is None:
                unmatched.append(f"{field_id}(field_id不存在)")
                continue
        else:
            label = item.get("label", "")
            data_key = label or f"第{i + 1}项"
            if not label:
                unmatched.append(f"第{i + 1}项(缺少field_id和label)")
                continue
            best_score = 0.0
            for f in form_fields:
                if f.get("field_id") in used_field_ids:
                    continue
                if f["label"] == label:
                    field = f
                    best_score = 1.0
                    break
                score = calculate_match_score(label, f["label"])
                if score > best_score:
                    best_score = score
                    field = f
            if field is None or best_score < 0.4:
                unmatched.append(label)
                continue

        used_field_ids.add(field.get("field_id"))
        matched.append({
            'field': field,
            'data_key': data_key,
            'field_label': field['label'],
            'value': value,
            'match_score': 1.0 if field_id else 0.9,
        })

    return matched, unmatched


def match_data_to_fields(
    data: Dict[str, Any],
    form_fields: List[Dict[str, Any]],
    manual_mapping: Dict[str, str],
) -> List[Dict[str, Any]]:
    """将数据键值对匹配到表单字段（自动匹配 + 手动映射优先）"""
    matched = []
    used_field_indexes = set()

    for data_key, field_label in manual_mapping.items():
        if data_key not in data:
            continue
        for idx, field in enumerate(form_fields):
            if idx in used_field_indexes:
                continue
            if field['label'] == field_label:
                matched.append({
                    'field': field, 'data_key': data_key, 'field_label': field_label,
                    'value': data[data_key], 'match_score': 1.0,
                })
                used_field_indexes.add(idx)

    for data_key, value in data.items():
        if any(m['data_key'] == data_key for m in matched):
            continue

        matching_fields = []
        for idx, field in enumerate(form_fields):
            if idx in used_field_indexes:
                continue
            score = calculate_match_score(data_key, field['label'])
            if score >= 0.4:
                matching_fields.append((idx, field, score))

        if not matching_fields:
            continue

        matching_fields.sort(key=lambda x: x[2], reverse=True)
        # 多个字段匹配同一标签时全部填充（如"组名"出现多次）
        best_score = matching_fields[0][2]
        for idx, field, score in matching_fields:
            if score >= best_score * 0.9:
                matched.append({
                    'field': field, 'data_key': data_key, 'field_label': field['label'],
                    'value': value, 'match_score': score,
                })
                used_field_indexes.add(idx)

    return matched


# ──────────────────────────── 纯文本变换（占位符拼接，不写文档） ────────────────────────────

def transform_underline(text: str, value: str, fill_mode: str) -> Optional[str]:
    """下划线占位符替换。append=在下划线前插入；overwrite=替换并保留余量下划线"""
    match = re.search(r'_{3,}', text)
    if not match:
        return None
    if fill_mode == 'append':
        return text[:match.start()] + value + " " + text[match.start():]
    underline_len = len(match.group())
    if len(value) >= underline_len:
        return text[:match.start()] + value + text[match.end():]
    remaining = '_' * (underline_len - len(value))
    return text[:match.start()] + value + remaining + text[match.end():]


_BRACKET_PATTERNS = [
    (re.compile(r'【\s*】'), '【', '】'),
    (re.compile(r'\[\s*\]'), '[', ']'),
    (re.compile(r'（\s*）'), '（', '）'),
]


def transform_bracket(text: str, value: str, fill_mode: str) -> Optional[str]:
    """方括号占位符替换（保留括号）"""
    for pattern, open_b, close_b in _BRACKET_PATTERNS:
        match = pattern.search(text)
        if match:
            if fill_mode == 'append':
                return text[:match.start()] + open_b + value + close_b + " " + text[match.end():]
            return text[:match.start()] + open_b + value + close_b + text[match.end():]
    return None


def transform_colon(text: str, value: str, fill_mode: str) -> Optional[str]:
    """行尾冒号后填写"""
    if not re.search(r'[：:]\s*$', text):
        return None
    if fill_mode == 'append':
        return text + value
    return text.rstrip() + value


def transform_paragraph_field(text: str, value: str, fill_mode: str) -> str:
    """LLM 检测的段落字段：依次尝试下划线/方括号/冒号，都没有则段末追加"""
    for transform in (transform_underline, transform_bracket, transform_colon):
        result = transform(text, value, fill_mode)
        if result is not None:
            return result
    return (text + value) if fill_mode == 'append' else (text.rstrip() + value)


# ──────────────────────────── 字段定位 + 写入准备 ────────────────────────────

def prepare_form_field(snapshot: DocSnapshot, field: Dict[str, Any],
                       value: Any, fill_mode: str) -> PreparedFill:
    """定位字段并准备写入（不写盘）"""
    location = field.get('location', {})
    placeholder_type = field.get('placeholder_type', '')
    label = field.get('label', '?')
    text_value = "" if value is None else str(value)

    # xlsx 纵向表单字段
    if placeholder_type == 'xlsx_cell':
        addr = XlsxCellAddr(sheet=location.get('sheet'), cell=location.get('cell', ''))
        return _prepare_xlsx_cell(snapshot, addr, field, value)

    # docx 内容控件
    if placeholder_type == 'content_control':
        addr = ContentControlAddr(tag=location.get('tag'), sdt_index=location.get('sdt_index'))
        return _prepare_content_control(snapshot, addr, text_value)

    # docx 表格单元格（含 llm_detected 的表格字段）
    if placeholder_type == 'table_cell' or (
        placeholder_type == 'llm_detected' and location.get('table_index') is not None
    ):
        addr = DocxTableCellAddr(
            table_index=location.get('table_index', -1),
            row=location.get('row_index', -1),
            col=location.get('cell_index', -1),
        )
        return _prepare_table_cell(snapshot, addr, text_value)

    # docx 段落字段（underline/bracket/colon/llm_detected 段落）
    transform_map = {
        'underline': transform_underline,
        'bracket': transform_bracket,
        'colon': transform_colon,
    }
    transform = transform_map.get(placeholder_type, transform_paragraph_field)
    addr = DocxParagraphAddr(
        paragraph_index=location.get('paragraph_index'),
        anchor=None,  # 结构分析时的段落文本可能已变，由 resolver 索引直达 + 报告展示
        index_scope="all",
    )
    return _prepare_paragraph(snapshot, addr, label, text_value, fill_mode, transform)


def _prepare_xlsx_cell(snapshot: DocSnapshot, addr: XlsxCellAddr,
                       field: Dict[str, Any], value: Any) -> PreparedFill:
    desc = describe_address(addr)
    resolved = AddressResolver.resolve(addr, snapshot)
    if not resolved.ok:
        return PreparedFill(subject=desc, ok=False, status=resolved.status, detail=resolved.detail)

    cell = resolved.target
    before = "" if cell.value is None else str(cell.value)
    # 数值类型字段按数值写入，避免 Excel 中数字变文本
    value_type = "number" if field.get('detected_type') == 'number' else None

    prepared = PreparedFill(subject=desc, ok=True, before=before, after="" if value is None else str(value))
    prepared.set_apply(lambda: ops.set_xlsx_cell(cell, value, value_type))
    return prepared


def _prepare_content_control(snapshot: DocSnapshot, addr: ContentControlAddr, value: str) -> PreparedFill:
    desc = describe_address(addr)
    resolved = AddressResolver.resolve(addr, snapshot)
    if not resolved.ok:
        return PreparedFill(subject=desc, ok=False, status=resolved.status, detail=resolved.detail)
    sdt = resolved.target

    from docx.oxml.ns import qn
    before = "".join(t.text or "" for t in sdt.findall('.//' + qn('w:t'))).strip()

    prepared = PreparedFill(subject=desc, ok=True, before=before, after=value)
    prepared.set_apply(lambda: ops.set_content_control_text(sdt, value))
    return prepared


def _prepare_table_cell(snapshot: DocSnapshot, addr: DocxTableCellAddr, value: str) -> PreparedFill:
    desc = describe_address(addr)
    resolved = AddressResolver.resolve(addr, snapshot)
    if not resolved.ok:
        return PreparedFill(subject=desc, ok=False, status=resolved.status, detail=resolved.detail)
    cell = resolved.target

    prepared = PreparedFill(subject=desc, ok=True, before=cell.text.strip(), after=value)
    prepared.set_apply(lambda: ops.set_cell_text(cell, value))
    return prepared


def _prepare_paragraph(snapshot: DocSnapshot, addr: DocxParagraphAddr, label: str,
                       value: str, fill_mode: str, transform) -> PreparedFill:
    desc = f"字段「{label}」{describe_address(addr)}"
    resolved = AddressResolver.resolve(addr, snapshot)
    if not resolved.ok:
        return PreparedFill(subject=desc, ok=False, status=resolved.status, detail=resolved.detail)
    para = resolved.target

    before = para.text.strip()
    new_text = transform(para.text, value, fill_mode)
    if new_text is None:
        return PreparedFill(
            subject=desc, ok=False, status="not_found",
            detail=f"段落中未找到对应占位符，当前内容: '{before[:50]}'",
        )

    prepared = PreparedFill(subject=desc, ok=True, before=before[:60], after=new_text.strip()[:60])
    prepared.set_apply(lambda: ops.set_paragraph_text(para, new_text))
    return prepared


def prepare_table_rows_field(snapshot: DocSnapshot, field: Dict[str, Any],
                             values: List[Dict[str, str]]) -> PreparedFill:
    """列表字段（is_header_based）：向表格表头下写入多行数据"""
    location = field.get('location', {})
    label = field.get('label', '?')
    table_idx = location.get('table_index')
    header_row_idx = location.get('header_row')

    desc = f"列表字段「{label}」表格{table_idx}"
    if table_idx is None or header_row_idx is None:
        return PreparedFill(subject=desc, ok=False, status="not_found", detail="缺少 table_index/header_row")

    doc = snapshot.docx
    if table_idx >= len(doc.tables):
        return PreparedFill(subject=desc, ok=False, status="not_found",
                            detail=f"表格索引 {table_idx} 超出范围（共 {len(doc.tables)} 个表格）")
    table = doc.tables[table_idx]
    if header_row_idx >= len(table.rows):
        return PreparedFill(subject=desc, ok=False, status="not_found",
                            detail=f"表头行 {header_row_idx} 超出范围（共 {len(table.rows)} 行）")

    headers = [cell.text.strip() for cell in table.rows[header_row_idx].cells]

    def _apply() -> ops.OpResult:
        for data_idx, data_row in enumerate(values):
            row_idx = header_row_idx + 1 + data_idx
            row = table.add_row() if row_idx >= len(table.rows) else table.rows[row_idx]
            for col_idx, header in enumerate(headers):
                if col_idx >= len(row.cells):
                    continue
                value = data_row.get(header, "")
                if not value:
                    for key, val in data_row.items():
                        if key in header or header in key:
                            value = val
                            break
                if value:
                    ops.set_cell_text(row.cells[col_idx], str(value))
        return ops.OpResult(ok=True, subject=desc, after=f"写入 {len(values)} 行")

    prepared = PreparedFill(subject=desc, ok=True, after=f"将写入 {len(values)} 行（表头行 {header_row_idx} 之下）")
    prepared.set_apply(_apply)
    return prepared
