"""
原子操作：对文档快照的封闭操作集

每个操作：
- 精确定位（调用方先用 AddressResolver 解析，或按句柄直接给目标对象）
- 执行
- 返回 OpResult{ok, before, after, error} —— before/after 供 LLM/用户验证

docx 文本写入统一"保留首 run 格式"策略（收编原 edit_cells_tool 的单元格写入
与 document_edit_tools 的 _replace_in_paragraph_runs 于本模块）。
"""

from typing import Any, List, Optional

from pydantic import BaseModel, Field


class OpResult(BaseModel):
    """原子操作结果"""
    ok: bool
    subject: str = Field("", description="操作对象描述（地址文本）")
    before: str = ""
    after: str = ""
    error: str = ""

    @classmethod
    def fail(cls, subject: str, error: str) -> "OpResult":
        return cls(ok=False, subject=subject, error=error)


# ──────────────────────────── docx 段落 ────────────────────────────

def set_paragraph_text(paragraph, new_text: str) -> OpResult:
    """设置段落文本，保留第一个 run 的格式"""
    before = paragraph.text
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = new_text
    return OpResult(ok=True, before=before, after=paragraph.text)


def replace_in_paragraph_runs(paragraph, old: str, new: str) -> int:
    """run 感知的段内文本替换，尽量保留混排格式；返回替换次数"""
    if not old or not paragraph.runs:
        return 0

    runs = paragraph.runs
    count = 0
    while True:
        full = "".join(r.text or "" for r in runs)
        idx = full.find(old)
        if idx < 0:
            break
        end = idx + len(old)

        pos = 0
        start_run = end_run = None
        start_off = end_off = 0
        for i, r in enumerate(runs):
            rlen = len(r.text or "")
            if start_run is None and idx < pos + rlen:
                start_run, start_off = i, idx - pos
            if start_run is not None and end <= pos + rlen:
                end_run, end_off = i, end - pos
                break
            pos += rlen

        if start_run is None or end_run is None:
            break

        if start_run == end_run:
            r = runs[start_run]
            text = r.text or ""
            r.text = text[:start_off] + new + text[end_off:]
        else:
            first = runs[start_run]
            last = runs[end_run]
            first.text = (first.text or "")[:start_off] + new
            for r in runs[start_run + 1:end_run]:
                r.text = ""
            last.text = (last.text or "")[end_off:]

        count += 1

    return count


def replace_in_paragraph(paragraph, old: str, new: str) -> OpResult:
    """段内替换（run 感知）；old 为空时整段替换"""
    before = paragraph.text
    if old:
        n = replace_in_paragraph_runs(paragraph, old, new)
        if n == 0:
            return OpResult.fail("段落", f"未找到文本: '{old[:50]}'，当前内容: '{before[:100]}'")
    else:
        set_paragraph_text(paragraph, new)
    return OpResult(ok=True, before=before, after=paragraph.text)


def insert_paragraph_after(paragraph, text: str = "", style: Optional[str] = None):
    """在指定段落后插入新段落并返回（python-docx 无原生 insert_paragraph_after）"""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    new_p = paragraph._p.makeelement(qn("w:p"), {})
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    if style and style != "Normal":
        try:
            new_para.style = style
        except (KeyError, ValueError):
            pass
    return new_para


def set_paragraph_style(paragraph, *, heading_level: Optional[int] = None,
                        list_type: Optional[str] = None,
                        font_name: Optional[str] = None,
                        font_size_pt: Optional[float] = None) -> OpResult:
    """设置段落样式（标题/列表/字体）"""
    before = f"[{paragraph.style.name if paragraph.style else 'Normal'}] {paragraph.text[:50]}"

    if heading_level is not None:
        level = max(1, min(6, int(heading_level)))
        paragraph.style = f"Heading {level}"
    if list_type:
        style_name = "List Number" if list_type == "number" else "List Bullet"
        try:
            paragraph.style = style_name
        except (KeyError, ValueError):
            pass
    if font_name or font_size_pt is not None:
        from docx.shared import Pt
        for run in paragraph.runs:
            if font_name:
                run.font.name = font_name
            if font_size_pt is not None:
                run.font.size = Pt(font_size_pt)

    after = f"[{paragraph.style.name if paragraph.style else 'Normal'}] {paragraph.text[:50]}"
    return OpResult(ok=True, before=before, after=after)


# ──────────────────────────── docx 表格 ────────────────────────────

def set_cell_text(cell, value: str, mode: str = "replace") -> OpResult:
    """写入表格单元格，保留第一个段落第一个 run 的格式"""
    before = cell.text
    new_text = before + value if mode == "append" else value

    if cell.paragraphs:
        first_para = cell.paragraphs[0]
        if first_para.runs:
            first_para.runs[0].text = new_text
            for run in first_para.runs[1:]:
                run.text = ""
        else:
            first_para.text = new_text
        for para in cell.paragraphs[1:]:
            for run in para.runs:
                run.text = ""
    else:
        cell.text = new_text

    return OpResult(ok=True, before=before, after=new_text)


def is_empty_row(row) -> bool:
    """表格行是否为空（所有单元格都为空或只有空白字符）"""
    if not row.cells:
        return True
    return all(not cell.text.strip() for cell in row.cells)


def delete_row(table, row) -> None:
    """从表格中删除一行（按行对象）"""
    table._tbl.remove(row._tr)


def append_row(table, values: List[str]):
    """在表格末尾追加一行并填入值（按列顺序，保留单元格格式）"""
    row = table.add_row()
    for i, value in enumerate(values):
        if i < len(row.cells):
            set_cell_text(row.cells[i], "" if value is None else str(value))
    return row


# ──────────────────────────── xlsx ────────────────────────────

def set_xlsx_cell(cell, value: Any, value_type: Optional[str] = None) -> OpResult:
    """写入 xlsx 单元格。value=None 表示清空；value_type='number' 强制数值"""
    before = cell.value

    if value is not None and value_type == "number":
        try:
            f = float(value)
            value = int(f) if f.is_integer() else f
        except (TypeError, ValueError):
            return OpResult.fail(cell.coordinate, f"值 '{value}' 无法转换为数值")

    cell.value = value
    return OpResult(
        ok=True,
        subject=cell.coordinate,
        before="" if before is None else str(before),
        after="" if value is None else str(value),
    )


# ──────────────────────────── 内容控件 ────────────────────────────

def set_content_control_text(sdt_element, value: str) -> OpResult:
    """写入内容控件：保留首个 w:r 的格式，清空其余文本节点"""
    from docx.oxml.ns import qn

    texts = sdt_element.findall('.//' + qn('w:t'))
    before = "".join(t.text or "" for t in texts)

    if texts:
        texts[0].text = value
        for t in texts[1:]:
            t.text = ""
    else:
        # 控件没有任何 run：在 w:sdtContent 下补一个最小 run
        content = sdt_element.find(qn('w:sdtContent'))
        if content is None:
            return OpResult.fail("内容控件", "控件缺少 sdtContent 节点")
        r = content.makeelement(qn('w:r'), {})
        t = content.makeelement(qn('w:t'), {})
        t.text = value
        r.append(t)
        content.append(r)

    return OpResult(ok=True, before=before, after=value)
